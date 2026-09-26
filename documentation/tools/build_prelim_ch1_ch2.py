from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import build_report as br


OUTPUT = br.DOC_DIR / "DELSU_Result_Advisory_Preliminaries_Chapters_1_2.docx"


def add_preliminaries(doc):
    title_style_ppr = doc.styles["Title"]._element.find(qn("w:pPr"))
    if title_style_ppr is not None:
        title_style_border = title_style_ppr.find(qn("w:pBdr"))
        if title_style_border is not None:
            title_style_ppr.remove(title_style_border)

    section = doc.sections[0]
    br.configure_section(section)
    section.different_first_page_header_footer = True
    section.footer.is_linked_to_previous = False
    br.set_page_number_start(section, 1, "lowerRoman")
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br.add_page_number(footer_p, roman=True)

    title = doc.add_paragraph(style="Title")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = 0
    title.add_run(br.TITLE)
    for _ in range(2):
        br.centred(doc, "")
    br.centred(doc, "BY", 12, True)
    br.centred(doc, br.AUTHOR, 14, True)
    br.centred(doc, br.MATRIC, 12, True)
    for _ in range(2):
        br.centred(doc, "")
    br.centred(doc, "A PROJECT REPORT SUBMITTED TO THE DEPARTMENT OF COMPUTER SCIENCE", 12, True)
    br.centred(doc, "FACULTY OF SCIENCE", 12, True)
    br.centred(doc, "DELTA STATE UNIVERSITY ABRAKA", 12, True)
    br.centred(doc, "IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE AWARD OF THE DEGREE OF", 12, True)
    br.centred(doc, "BACHELOR OF SCIENCE IN COMPUTER SCIENCE", 12, True)
    for _ in range(2):
        br.centred(doc, "")
    br.centred(doc, "SEPTEMBER 2026", 12, True)

    br.heading(doc, "DECLARATION", 1)
    br.body(doc, "I declare that this project report is an account of the system I designed and developed. All ideas and materials obtained from other sources have been acknowledged through in-text citations and the reference list. This work has not been submitted in full for another degree or academic award.")
    signature = br.body(doc, "Name: Dennis Emmanuel\nMatriculation Number: FOS/22/23/292155\nSignature: ______________________________\nDate: __________________________________", indent=False)
    signature.alignment = WD_ALIGN_PARAGRAPH.LEFT
    signature.paragraph_format.line_spacing = 1.5

    br.heading(doc, "CERTIFICATION", 1)
    br.body(doc, "This project report titled Development of an AI Powered Student Academic Performance Analysis and Advisory System was carried out by Dennis Emmanuel, with matriculation number FOS/22/23/292155, in the Department of Computer Science, Delta State University, Abraka. It is submitted in partial fulfilment of the requirements for the award of the Bachelor of Science degree in Computer Science.")
    br.table(doc, "P.1", "Certification signatures", ["Role", "Name", "Signature", "Date"], [
        ["Project supervisor", "", "", ""],
        ["Head of department", "", "", ""],
        ["External examiner", "", "", ""],
    ], [1.6, 2.0, 1.5, 1.1])

    br.heading(doc, "DEDICATION", 1)
    br.body(doc, "This project is dedicated to my family, teachers and classmates for their support, encouragement and contribution to my education.")

    br.heading(doc, "ACKNOWLEDGEMENTS", 1)
    br.body(doc, "I give thanks to God for the strength and wisdom to complete this project. I appreciate my project supervisor for the guidance, corrections and encouragement given throughout the work. I also thank the lecturers in the Department of Computer Science, Delta State University, Abraka, for the knowledge and training I received during my undergraduate programme. I am grateful to the course advisers who explained the result preparation process and the difficulties involved in distributing academic records to students. Finally, I thank my family, friends and classmates for their patience, assistance and encouragement during the development of the system.")

    br.heading(doc, "ABSTRACT", 1)
    br.body(doc, "Academic results provide important information about a student's progress, but the method used to distribute them can affect privacy, accessibility and understanding. In the Department of Computer Science, result broadsheets are prepared with Microsoft Excel and may be shared as large static files when students need access. This approach requires each student to search for a record manually and does not provide personal analysis or immediate academic guidance. This project developed Compass, an AI powered student academic performance analysis and advisory system. The system provides separate interfaces for students and course advisers. Course advisers can upload and preview result broadsheets, while students can view their results, GPA, CGPA, outstanding courses and notifications through authenticated accounts. The advisory component uses defined backend tools to retrieve stored academic information and explain it in simple language. The system was implemented with React, FastAPI, Supabase Postgres, Supabase Auth and a hosted large language model. An iterative development method was used so that result processing, user interfaces and advisory functions could be developed and checked in stages. The completed prototype demonstrates how an existing spreadsheet workflow can be connected to a role based web application that gives each student a private and interactive view of academic performance. The project therefore provides a practical foundation for improving result access and academic support within the department.")

    br.heading(doc, "TABLE OF CONTENTS", 1)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = 0
    br.add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Update the table of contents in Microsoft Word.")

    br.heading(doc, "LIST OF TABLES", 1)
    br.body(doc, "Table 2.1  Comparison of related approaches", indent=False)

    br.heading(doc, "LIST OF FIGURES", 1)
    br.body(doc, "No figures are included in this document.", indent=False)

    br.heading(doc, "LIST OF ABBREVIATIONS", 1)
    br.table(doc, "P.2", "Abbreviations used in the report", ["Abbreviation", "Meaning"], [
        ["AI", "Artificial intelligence"],
        ["API", "Application programming interface"],
        ["CGPA", "Cumulative grade point average"],
        ["DELSU", "Delta State University"],
        ["GPA", "Grade point average"],
        ["ICT", "Information and communication technology"],
        ["LLM", "Large language model"],
        ["RBAC", "Role based access control"],
        ["REST", "Representational State Transfer"],
        ["RLS", "Row level security"],
    ], [1.7, 4.5])


def chapter_one(doc):
    br.chapter(doc, "ONE", "INTRODUCTION")

    br.heading(doc, "1.1 Background to the study", 2)
    br.body(doc, "Universities depend on academic records to monitor student progress, determine eligibility for graduation and provide guidance. These records normally contain course scores, grades, credit units, grade points, semester GPA and cumulative GPA. When they are accurate and easy to access, students can understand their academic position and make better decisions about their studies. Course advisers also use the same information to identify outstanding courses, confirm progression requirements and provide support where necessary.")
    br.body(doc, "The use of information and communication technology has changed the way educational institutions store and present academic information. Paper registers and isolated spreadsheets are gradually being replaced by database systems and web portals. A web based system can make information available through a browser, separate the records of different users and present calculated summaries without requiring a student to search through a large file. In Nigerian higher education, however, institutions continue to face challenges involving infrastructure, system availability, maintenance and the effective use of digital platforms. These challenges can cause departments to depend on familiar tools such as Microsoft Excel when preparing and sharing results.")
    br.body(doc, "Microsoft Excel is useful for result preparation because it supports tabular data, formulas and familiar checking procedures. A course adviser can enter continuous assessment scores, examination scores and course units in a broadsheet, then calculate grades and summary values. The limitation appears when the same broad class file is distributed to students. A student must search for a matriculation number in a large document, and the file may contain the records of many other students. The method also provides little assistance in explaining GPA, CGPA, carryovers or the possible effect of future grades.")
    br.body(doc, "Learning analytics provides a useful direction for improving this process. It involves collecting and analysing educational data to understand learning and support decisions. Studies have shown that academic information such as previous grades, attendance and student activity can help institutions identify patterns and provide early support (Namoun & Alshanqiti, 2021; Albreiki et al., 2021). The present project does not attempt to replace official academic judgement with automatic prediction. Instead, it uses verified result records to produce summaries that students and advisers can understand.")
    br.body(doc, "Artificial intelligence also creates new ways of interacting with information systems. Modern large language models can receive a question in ordinary language and produce a relevant response. Their value is greater when they can use external tools instead of depending only on information learned during training. ReAct combines reasoning with actions in an external environment, while Toolformer and later studies explain how language models can call tools and use returned observations (Yao et al., 2023; Schick et al., 2023; Qin et al., 2023). In an academic result system, a tool can retrieve a student's stored GPA, CGPA or course record before the model prepares an explanation.")
    br.body(doc, "This project applies these ideas to the Department of Computer Science at Delta State University, Abraka. The developed system is called Compass. It accepts the department's existing Excel result format, processes the records and stores them in a structured database. Students sign in to view their personal academic information, while course advisers use a separate interface for uploads and class summaries. Compass also allows a student to ask questions about academic performance and receive an answer based on stored records. The project therefore combines result management, learning analytics, role based access and tool supported conversational assistance in one web application.")

    br.heading(doc, "1.2 Statement of the problem", 2)
    br.body(doc, "Students in the Department of Computer Science need a reliable and private way to access their academic results. The current spreadsheet based process is effective for preparing a class broadsheet, but the distribution method is not designed for personal access. When a broad result file is shared, each student receives more information than is required and must search through many rows to find one record. This makes the process inconvenient, especially on a mobile device, and may expose the academic details of other students.")
    br.body(doc, "The static file also gives limited support for understanding academic performance. It may show course scores and calculated totals, but it does not explain how GPA or CGPA was obtained, identify outstanding compulsory courses in a personal view or allow a student to test the effect of possible future grades. Students may therefore depend on repeated explanations from course advisers for questions that can be answered from the available result data.")
    br.body(doc, "Course advisers face a related challenge. They work with large broadsheets that contain many students, courses, units and formulas. Before a result is distributed, the file should be checked for missing values, invalid scores and differences between supplied totals and recalculated values. A static sharing process does not provide a separate preview and confirmation stage, and it does not create a structured history of result uploads and student notifications.")
    br.body(doc, "There is therefore a need for a system that preserves the familiar Excel input process while improving result validation, personal access and academic guidance. The required system should allow course advisers to upload and review result data, store the confirmed records in a database and give each authenticated student access to an individual academic profile. It should also provide simple explanations based on the student's verified information rather than producing unsupported academic values.")

    br.heading(doc, "1.3 Aim and objectives", 2)
    br.body(doc, "The aim of this project is to develop an AI powered student academic performance analysis and advisory system for the Department of Computer Science at Delta State University.")
    br.body(doc, "The specific objectives are to:", indent=False)
    br.numbered_list(doc, [
        "design a role based web platform that gives students personal access to their academic results and provides course advisers with an administrative workspace;",
        "develop a result ingestion and validation process for departmental Excel broadsheets;",
        "implement an academic advisory agent that retrieves stored student records and supports grade simulations;",
        "provide academic summaries such as GPA, CGPA, degree classification and outstanding courses; and",
        "evaluate the completed system through build checks, functional testing, API testing and database inspection.",
    ])

    br.heading(doc, "1.4 Research questions", 2)
    br.body(doc, "The project is guided by the following research questions:", indent=False)
    br.numbered_list(doc, [
        "How can departmental result broadsheets be converted into private student records without removing the course adviser's familiar Excel workflow?",
        "How can the system identify invalid entries and calculation differences before result data is confirmed?",
        "How can an AI advisory component answer student questions with information retrieved from verified academic records?",
        "How can the system present GPA, CGPA, outstanding courses and class summaries in a form that is easy to understand?",
        "What do functional, performance and database tests show about the operation of the developed prototype?",
    ])

    br.heading(doc, "1.5 Significance of the study", 2)
    br.body(doc, "The project is important to students because it provides a personal and convenient way to view academic information. A student can sign in and see course results, GPA, CGPA, degree classification and outstanding courses without searching through a class broadsheet. The information is arranged in a focused interface that can be used on desktop and mobile screens. Notifications can also inform the student when new result data becomes available.")
    br.body(doc, "The advisory feature adds an interactive form of support. A student can ask a question about current academic standing or the effect of possible grades. The system retrieves the relevant stored values and presents the explanation in simple language. This can help students understand their records before meeting a course adviser, while final academic decisions remain the responsibility of authorised university staff.")
    br.body(doc, "Course advisers benefit from a structured upload process. The system accepts an Excel file, processes its rows and produces a preview before confirmation. This provides an opportunity to identify invalid values and inspect calculated outcomes. The adviser interface also presents class summaries such as the number of students, average CGPA, degree class distribution and students who may require attention.")
    br.body(doc, "The department benefits from a central database that separates data storage from file distribution. Confirmed results can support student views, adviser summaries, notifications and advisory questions without creating different copies of the same academic record. The project also demonstrates how modern web development and AI tool use can be applied to a practical problem in a Nigerian university environment.")
    br.body(doc, "The study contributes to computer science practice by combining frontend development, backend services, database design, authentication, spreadsheet processing and language model tool calling. It provides a working example of how an AI component can be connected to a conventional information system while calculations and official values remain in defined application functions.")

    br.heading(doc, "1.6 Scope of the study", 2)
    br.body(doc, "The project covers the management and presentation of academic results for students in the Department of Computer Science at Delta State University. The system processes Microsoft Excel broadsheets in supported wide and long layouts. It stores student, course and result information, calculates semester GPA and cumulative GPA, identifies failed compulsory courses and presents academic summaries.")
    br.body(doc, "The main user roles are student, course adviser and administrator. Students can authenticate, view their academic records, read notifications and use the Compass advisory interface. Course advisers can authenticate, upload and preview broadsheets, confirm valid result data, view upload history and access class analytics. The administrator role supports the approval and management of adviser accounts.")
    br.body(doc, "The advisory component is limited to questions and calculations supported by the available backend tools. It can retrieve academic records, explain performance and calculate possible outcomes from values supplied by a student. It does not approve results, change grades, register courses, pay fees or make final decisions about graduation. These responsibilities remain with the university and its authorised officers.")
    br.body(doc, "The project focuses on a departmental prototype. It does not replace the official university result system or the approved senate process. Its purpose is to demonstrate a practical method for improving access, analysis and guidance around verified result data.")

    br.heading(doc, "1.7 Limitations of the study", 2)
    br.body(doc, "The project was developed and evaluated as a prototype with available departmental result data and selected user accounts. It did not include a large controlled usability study involving a representative sample of students and advisers. User satisfaction, learning outcomes and long term changes in adviser workload were therefore not measured.")
    br.body(doc, "The system depends on internet access and external services for authentication, database access and language model responses. Network conditions can affect loading and response time. The supported spreadsheet formats are based on the broadsheets examined during development, so a file with a different structure may require additional mapping rules.")
    br.body(doc, "The advisory response is limited by the quality of the stored data, the available tools and the clarity of the student's question. Although the system retrieves verified academic values, its explanations are intended to support understanding and do not replace advice from qualified staff. The project also does not cover course registration, fee payment, examination scheduling or the full range of university student services.")

    br.heading(doc, "1.8 Definition of terms", 2)
    terms = [
        ("Academic advisory agent", "A software component that retrieves academic information and produces explanations or calculations within defined limits."),
        ("Broadsheet", "An Excel worksheet that contains the scores, grades and calculated results of several students."),
        ("Carryover", "A failed course that a student is required to retake according to the applicable academic rules."),
        ("CGPA", "The cumulative grade point average calculated from grade points and credit units across the academic periods represented in the system."),
        ("GPA", "The weighted average of grade points earned during a defined semester."),
        ("Learning analytics", "The collection and analysis of educational data to understand learning and support decisions."),
        ("Large language model", "An artificial intelligence model trained on large amounts of text to understand instructions and generate language."),
        ("Role based access control", "A method of assigning system permissions according to the responsibilities of a user role."),
        ("Row level security", "A database control that determines which records a user may read or change."),
        ("Tool calling", "A process in which a language model requests a defined software function and uses the returned information in its answer."),
    ]
    for term, meaning in terms:
        p = br.body(doc, "", indent=False)
        p.add_run(term + ": ").bold = True
        p.add_run(meaning)


def chapter_two(doc):
    br.chapter(doc, "TWO", "LITERATURE REVIEW")
    br.body(doc, "This chapter reviews the concepts, theories and previous studies that support the development of the proposed system. It covers academic result management, learning analytics, large language models, tool supported AI agents, conversational support, privacy, access control and explainability. The chapter also presents the theoretical framework, reviews related works and identifies the gap addressed by the project.")

    br.heading(doc, "2.1 Conceptual review", 2)
    br.heading(doc, "2.1.1 Academic result management systems", 3)
    br.body(doc, "An academic result management system receives assessment data, applies grading rules, stores the outcome and presents information for later use. Its main activities include data entry, validation, calculation, storage, retrieval and reporting. These activities may be completed with paper records, spreadsheets or database applications. The selected method affects accuracy, accessibility, privacy and the amount of work required when a correction is made.")
    br.body(doc, "Spreadsheets remain common because they provide a familiar table structure and allow users to create formulas without building a complete application. In a result broadsheet, each row can represent a student while course scores and calculated values appear across columns. The same file can support checking and printing. However, a spreadsheet does not provide personal access when the complete file is distributed. It also becomes difficult to maintain when several copies are created after corrections.")
    br.body(doc, "A database system stores records separately from the interface used to view them. This makes it possible to show each student only the information connected to an authenticated account. It also allows the same stored data to support several functions, including student result pages, adviser summaries, notifications and reports. Database relationships can connect students, courses, uploads and individual result rows so that the origin of each value can be identified.")
    br.body(doc, "Web based result systems improve accessibility because users can reach them through a browser. Their quality still depends on correct calculations, clear interfaces, reliable services and appropriate access controls. ISO/IEC 25010:2023 groups software quality around characteristics such as functional suitability, performance efficiency, usability, reliability, maintainability and security. These characteristics provide a useful basis for thinking about an academic result system as more than a calculation tool.")

    br.heading(doc, "2.1.2 Learning analytics and student performance", 3)
    br.body(doc, "Learning analytics uses data about learners and their activities to understand and improve education. Academic records are an important source because they show performance across courses and periods. Institutions can use summaries such as average scores, grade distributions and progression rates to understand class performance. Students can also benefit when the same information is presented in a personal form that explains current standing.")
    br.body(doc, "Namoun and Alshanqiti (2021) reviewed data mining and learning analytics studies concerned with student performance. Their review showed frequent use of prior grades, demographic information, attendance and learning activity. Albreiki et al. (2021) similarly examined machine learning methods for performance prediction. These studies show that academic data can support early identification and intervention, although the method and available data affect the meaning of any prediction.")
    br.body(doc, "Performance prediction is one part of learning analytics, but not every useful system requires a predictive model. Direct calculations can provide clear information about GPA, CGPA, failed courses and degree classification. A student can understand how a value was obtained when it is linked to course units and grade points. The present project uses this direct approach for its main summaries and uses the advisory component to explain the stored values rather than predict an official result.")
    br.body(doc, "Dashboards are commonly used to present learning analytics because they place important measures in one interface. A student dashboard may show personal progress, while an adviser dashboard may show class totals, averages and distributions. The value of a dashboard depends on selecting information that supports a real task. Too many measures can make the interface confusing, while a small set of clear values can help the user decide what to examine next.")

    br.heading(doc, "2.1.3 Large language models", 3)
    br.body(doc, "Large language models are computer models trained to process and generate text. Their modern development is closely linked to the Transformer architecture introduced by Vaswani et al. (2017). The Transformer uses attention mechanisms to identify relationships among words in a sequence. This approach made it possible to train larger models that can follow instructions, answer questions and perform several language tasks.")
    br.body(doc, "BERT demonstrated the value of pretraining for language understanding tasks (Devlin et al., 2019). Brown et al. (2020) later showed that a sufficiently large language model could perform new tasks from instructions and a small number of examples. Libraries such as Transformers also made pretrained models more accessible to researchers and developers (Wolf et al., 2020). Zhao et al. (2023) reviewed the development, abilities and challenges of large language models across many applications.")
    br.body(doc, "A language model generates a response from patterns learned during training and the information supplied in the current prompt. It does not automatically have an accurate copy of a university database. This distinction is important for academic advising because GPA, CGPA and course status must agree with official records. Ji et al. (2023) describe hallucination as generated content that is not supported by the required source. Connecting the model to defined retrieval and calculation tools provides a practical way to reduce this problem.")
    br.body(doc, "The language model in Compass is used as an explanation layer. Academic calculations remain in application functions, while the model receives the values needed for the student's question. This separation allows the conversational interface to remain flexible without making the model the source of official academic data.")

    br.heading(doc, "2.1.4 Tool supported and agentic AI", 3)
    br.body(doc, "An AI agent combines language processing with actions that affect or inspect an external environment. The action may involve searching a document, calling an API, using a calculator or retrieving a database record. Tool use extends the model beyond the information contained in a prompt and allows an application to provide current observations.")
    br.body(doc, "ReAct joins reasoning and acting in a repeated process (Yao et al., 2023). A model interprets the task, selects an action, receives an observation and uses the result to continue. Toolformer explored how language models can learn to call external tools for information and computation (Schick et al., 2023). Qin et al. (2023) organised tool learning around task planning, tool selection, tool use and response generation. Together, these works provide a foundation for an advisory system that retrieves academic records before answering.")
    br.body(doc, "Chain of thought prompting showed that intermediate reasoning steps can improve performance on some tasks (Wei et al., 2022). In an operational information system, however, reasoning alone cannot supply a missing institutional fact. A reliable process must obtain the required value from an authorised source. For example, a question about current CGPA should call the relevant academic function instead of asking the model to estimate the value.")
    br.body(doc, "Agent evaluation should examine the complete sequence of actions, observations and responses. AgentBench studies language models in environments where they must act rather than produce only a single text answer (Liu et al., 2024). This view is relevant to Compass because the quality of an answer depends on the selected tool, the returned record, the calculation and the final explanation.")

    br.heading(doc, "2.1.5 Conversational support in education", 3)
    br.body(doc, "Conversational systems allow users to request information in ordinary language. In education, they can answer frequently asked questions, guide learners to resources and provide explanations outside normal office hours. Their availability can make support easier to reach, especially when a student needs a simple explanation of information that already exists in an institutional system.")
    br.body(doc, "Research on conversational AI in education reports both opportunities and responsibilities. Lo (2023) reviewed early studies on ChatGPT in education and identified possible uses for teaching, learning and support. Kasneci et al. (2023) discussed personalised explanation and assistance, together with concerns involving incorrect information, bias and inappropriate dependence. Tlili et al. (2023) also examined educational uses and ethical questions through a case based analysis.")
    br.body(doc, "A student result adviser should have a clear purpose. It can explain a calculated value, identify courses in the stored record and perform defined simulations. It should not present itself as a replacement for lecturers, course advisers or university authorities. UNESCO (2023) recommends human oversight, protection of learner data and clear responsibility when generative AI is used in education.")
    br.body(doc, "Compass follows a limited academic role. The student can ask about results, GPA, CGPA, outstanding courses and possible grade outcomes. The system retrieves relevant data through backend tools and presents an explanation. Questions outside the available academic functions should be redirected to the appropriate university office or qualified member of staff.")

    br.heading(doc, "2.1.6 Privacy and access control", 3)
    br.body(doc, "Academic results are personal records. A system that stores or presents them should limit access to users with a valid purpose. The Nigeria Data Protection Act 2023 provides a national framework for the lawful processing, security and rights associated with personal data. In an academic system, these principles support the need to collect only required information, protect account credentials and prevent one student from viewing another student's record.")
    br.body(doc, "Authentication and authorisation perform different functions. Authentication confirms the identity of a user, while authorisation determines the actions and records available to that identity. A valid sign in is therefore only the first step. The application must also decide whether the account belongs to a student, adviser or administrator and apply the permissions of that role.")
    br.body(doc, "Role based access control assigns permissions according to job responsibilities (Sandhu et al., 1996; Ferraiolo et al., 2001). Students require access to their own results and conversations. Course advisers require broader access for assigned classes, result uploads and summaries. Administrators require account management functions. These roles should be reflected in the interface, backend routes and database rules.")
    br.body(doc, "Postgres row level security provides another control at the database level. It can restrict a query according to the authenticated user or a relationship stored in the database. Supabase combines Postgres with authentication services, making it possible to link an account identifier to a student or adviser profile. Applying the same ownership rules across the client, application server and database helps maintain consistent access decisions.")

    br.heading(doc, "2.1.7 Explainability in educational systems", 3)
    br.body(doc, "Explainability concerns the ability of a user to understand how a result or recommendation was produced. It is important in education because academic information can influence confidence, course planning and requests for support. Khosravi et al. (2022) explain that educational AI should connect technical explanations with the needs of students and educators.")
    br.body(doc, "Some academic explanations can be produced directly from formulas. GPA can be explained through course units, grade points and quality points. CGPA can be explained through cumulative units and points. These calculations are easier to check than a prediction produced by an unknown internal pattern. A system should therefore present the evidence that matters to the user's question.")
    br.body(doc, "Conversational explanation adds another layer. The system can translate calculated values into ordinary language and identify the courses connected to them. This does not require exposing hidden model reasoning. It requires showing the verified academic inputs, the applied rule and the resulting value in a form the student can understand.")

    br.heading(doc, "2.1.8 Responsible use of generative AI", 3)
    br.body(doc, "Generative AI should be used with clear boundaries when it is connected to student records. The response may sound confident even when its evidence is incomplete. NIST (2023) treats validity, reliability, privacy, transparency and accountability as connected parts of AI risk management. These ideas are relevant to educational systems because students may act on information supplied by the application.")
    br.body(doc, "The responsible design used in this project keeps academic calculations in defined functions and limits the advisory tools to permitted tasks. The model does not receive unrestricted access to the database and cannot change an official result. It receives selected observations required for an explanation or simulation. Final decisions about grades, progression and graduation remain with authorised university officers.")
    br.body(doc, "The system should also make its purpose clear to the user. A student should understand that Compass provides academic information and guidance based on available records. Where the question requires an official decision or information outside the stored data, the student should contact a course adviser or the appropriate university office.")

    br.heading(doc, "2.2 Theoretical framework", 2)
    br.heading(doc, "2.2.1 ReAct framework", 3)
    br.body(doc, "The main theoretical framework for the advisory component is ReAct. The framework combines reasoning and action so that a language model can interact with an environment before producing a final response (Yao et al., 2023). In Compass, a student's question begins a cycle in which the model identifies the required information, selects an academic tool, receives the returned value and prepares an explanation.")
    br.body(doc, "A question about academic standing may require the cumulative GPA tool, while a question about outstanding courses requires the relevant course record. A grade simulation requires a calculation function that accepts proposed scores or grades. The tool output becomes the observation used in the answer. This makes the database and application functions the source of academic facts.")
    br.body(doc, "The ReAct framework is suitable for this project because the advisory task involves both language and information retrieval. A purely conversational model could explain the meaning of GPA, but it could not know the student's current value without access to the record. The action step provides that access through a controlled interface.")

    br.heading(doc, "2.2.2 Role based access control model", 3)
    br.body(doc, "Role based access control is the security framework for separating student, adviser and administrator functions. A role represents a set of responsibilities, while permissions define the operations allowed for that role. This model is appropriate for an academic result system because users have different duties and should not receive the same level of access.")
    br.body(doc, "The student role is mainly personal and read oriented. It permits access to the student's profile, academic record, notifications and advisory conversations. The adviser role permits result upload, preview, confirmation and class analysis within an assigned scope. The administrator role manages adviser approval and access status. The model guides navigation, API checks and database policies.")
    br.body(doc, "RBAC also supports maintainability. When permissions are connected to roles, the system can apply a consistent rule across several pages and services. A change in responsibility can be reflected in the role definition instead of being implemented separately in every interface component.")

    br.heading(doc, "2.2.3 Client server and REST architecture", 3)
    br.body(doc, "The project also follows a client server model. The React frontend is the client that presents pages and receives user input. The FastAPI backend processes requests, applies academic rules and communicates with the database. Supabase provides data storage and authentication, while the language model service supports the advisory response.")
    br.body(doc, "REST describes an architectural style in which a client communicates with server resources through a uniform interface (Fielding, 2000). In the developed system, HTTP routes represent profiles, students, results, uploads, analytics, notifications and conversations. The frontend sends a request and the backend returns structured JSON data or a streamed response.")
    br.body(doc, "This separation supports clear responsibilities. The frontend focuses on presentation and interaction, while the backend performs validation, calculations and controlled data access. The same backend function can serve a result page or an advisory tool without placing academic formulas in the browser.")

    br.heading(doc, "2.3 Review of related works", 2)
    br.body(doc, "Namoun and Alshanqiti (2021) conducted a systematic review of student performance prediction using data mining and learning analytics. The reviewed studies used academic and behavioural features to identify patterns associated with performance. The work demonstrates the value of educational data, but its main focus is prediction rather than private result delivery and conversational explanation.")
    br.body(doc, "Albreiki et al. (2021) reviewed machine learning techniques used for student performance prediction. They compared algorithms, datasets and evaluation measures across existing studies. Their findings support careful selection of methods and clear reporting of evaluation results. The present project differs by using direct academic calculations for its main student summaries instead of training a prediction model.")
    br.body(doc, "Kabathova and Drlik (2021) examined dropout prediction in university courses with several machine learning methods. The study shows how course data can support early identification of students who may need assistance. Its output is mainly intended for institutional monitoring and does not provide an individual student with a tool supported conversation about verified result records.")
    br.body(doc, "Zawacki-Richter et al. (2019) reviewed artificial intelligence applications in higher education. They identified common areas such as profiling, prediction, assessment and adaptive systems. The review also observed that educators were not always strongly represented in the research. This supports the need to design educational applications around the practical work of students and staff.")
    br.body(doc, "Ouyang et al. (2022) reviewed empirical research on artificial intelligence in online higher education. The study grouped applications around prediction, recommendation, assessment and support. It shows the broad role of AI in education and the importance of connecting a selected technique to a clear educational purpose. Compass uses AI for explanation and interaction while the academic record remains in a conventional database.")
    br.body(doc, "Kasneci et al. (2023) discussed opportunities and challenges associated with large language models in education. The authors identified possible value in personalised support and explanation, together with risks involving inaccurate information and bias. Their work supports an application design in which the language model has a limited role and important values are supplied by verified tools.")
    br.body(doc, "Khosravi et al. (2022) reviewed explainable artificial intelligence in education. They showed that explanation needs differ among students, instructors and administrators. This is relevant to the separation of interfaces in Compass. Students receive personal records and simple explanations, while course advisers receive class summaries that support academic follow up.")
    br.body(doc, "Schick et al. (2023), Yao et al. (2023) and Qin et al. (2023) examined different aspects of language model tool use. Their work shows that a model can use external functions to obtain information required for a task. The present project applies this idea to student result data by providing academic retrieval and simulation functions through the backend.")
    br.body(doc, "Park et al. (2023) explored generative agents that maintain observations and memory across interactions. Their work demonstrates how stored experience can support continuity in a conversation. Compass uses a simpler form of memory through saved chat sessions and messages so that a student can return to previous academic questions.")
    br.body(doc, "The related works show progress in learning analytics, predictive systems, conversational AI, explainability and tool supported agents. They also show that these areas are often studied separately. Performance systems may focus on prediction for administrators, while conversational systems may not have controlled access to verified institutional records.")

    br.table(doc, "2.1", "Comparison of related approaches", ["Approach", "Personal access", "Verified records", "Interactive guidance", "Spreadsheet upload", "Main emphasis"], [
        ["Shared result file", "No", "Static file", "No", "Not applicable", "Distribution of a broad class record"],
        ["Standard student portal", "Yes", "Yes", "Limited", "Varies", "Personal result viewing"],
        ["Performance prediction system", "Usually limited", "Dataset based", "No", "Varies", "Prediction and early identification"],
        ["General educational chatbot", "Yes", "Usually no", "Yes", "No", "Conversation and general support"],
        ["Compass", "Yes", "Yes through backend tools", "Yes", "Yes", "Integrated result access, analysis and guidance"],
    ], [1.2, 0.85, 1.1, 1.0, 1.0, 1.6], page_break_before=True)

    br.heading(doc, "2.4 Research gap", 2)
    br.body(doc, "The literature shows strong interest in student performance analysis and the use of conversational AI in education. Performance studies often apply machine learning to identify patterns or predict outcomes. Student portals provide personal access to results, while conversational systems provide natural language interaction. However, these functions are not always combined in a way that uses verified institutional records.")
    br.body(doc, "A second gap concerns the connection between existing departmental work and a new information system. Course advisers already prepare results in structured Excel broadsheets. Replacing that process completely may create additional work and reduce acceptance. There is a need for a system that accepts the familiar input format, validates it and converts the confirmed rows into individual student records.")
    br.body(doc, "A third gap concerns academic explanations. A standard portal may display GPA and course results without explaining them, while a general chatbot may produce an answer without access to the student's official data. A useful advisory system should combine personal access with controlled tools that retrieve the relevant record before an explanation is produced.")
    br.body(doc, "Compass addresses these gaps by integrating Excel result ingestion, structured database storage, role based interfaces, academic summaries, notifications and a tool supported advisory agent. The contribution is the combination of these functions in a departmental prototype designed around the result preparation process used in the selected environment.")

    br.heading(doc, "2.5 Summary of the literature review", 2)
    br.body(doc, "This chapter reviewed the concepts and studies that support the project. Academic result systems require accurate calculation, structured storage and suitable presentation. Learning analytics explains how result data can be converted into useful summaries, while large language model research explains the conversational abilities used in the advisory interface.")
    br.body(doc, "Tool learning and the ReAct framework provide the basis for retrieving academic information before producing an answer. RBAC and row level security support the separation of student, adviser and administrator access. Explainability and responsible AI principles show the importance of presenting evidence, limiting the role of the model and keeping final authority with qualified university staff.")
    br.body(doc, "The review of related works identified a gap between result portals, performance analysis systems and conversational educational support. The proposed system responds to this gap by connecting the existing spreadsheet workflow to a private web platform and a tool supported academic adviser. The next chapter of the complete project report will describe the analysis and design of the system.")


def add_references(doc):
    br.add_references(doc)


def build():
    doc = Document()
    br.setup_styles(doc)
    add_preliminaries(doc)

    main_section = doc.add_section(WD_SECTION.NEW_PAGE)
    br.configure_section(main_section)
    main_section.different_first_page_header_footer = False
    main_section.header.is_linked_to_previous = False
    main_section.footer.is_linked_to_previous = False
    br.set_page_number_start(main_section, 1, "decimal")
    header_p = main_section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    br.add_page_number(header_p, roman=False)
    main_section.footer.paragraphs[0].text = ""

    chapter_one(doc)
    chapter_two(doc)
    add_references(doc)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    props = doc.core_properties
    props.title = br.TITLE.title()
    props.subject = "Preliminary pages and Chapters One and Two"
    props.author = br.AUTHOR.title()
    props.keywords = "academic results, student advisory system, learning analytics, role based access, agentic AI"

    br.DOC_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
