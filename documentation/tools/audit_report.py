import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    full_text = "\n".join(paragraphs)

    abstract_start = paragraphs.index("ABSTRACT") + 1
    abstract_end = paragraphs.index("TABLE OF CONTENTS")
    abstract_text = " ".join(p for p in paragraphs[abstract_start:abstract_end] if p)

    references_start = paragraphs.index("REFERENCES") + 1
    references_end = paragraphs.index("APPENDICES")
    references = [p for p in paragraphs[references_start:references_end] if p]

    with zipfile.ZipFile(path) as archive:
        xml_text = "".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
        media_count = sum(1 for name in archive.namelist() if name.startswith("word/media/"))

    humanizer_phrases = {
        phrase: len(re.findall(rf"\b{re.escape(phrase)}\b", full_text, flags=re.IGNORECASE))
        for phrase in ("delve", "pivotal", "crucial", "stands as", "serves as", "not only", "not just")
    }
    report = {
        "abstract_words": len(re.findall(r"\b[\w']+\b", abstract_text)),
        "reference_entries": len(references),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "embedded_media": media_count,
        "em_dash_count": xml_text.count("\u2014"),
        "en_dash_count": xml_text.count("\u2013"),
        "placeholder_terms": {
            term: full_text.lower().count(term)
            for term in ("todo", "tbd", "right click and update", "insert figure", "lorem ipsum")
        },
        "humanizer_phrase_counts": humanizer_phrases,
    }
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
