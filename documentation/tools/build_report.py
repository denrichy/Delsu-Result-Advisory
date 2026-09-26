from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOC_DIR = ROOT / "documentation"
SCREEN = DOC_DIR / "assets" / "screenshots"
DIAGRAM = DOC_DIR / "assets" / "diagrams"
OUTPUT = DOC_DIR / "DELSU_Result_Advisory_Final_Project_Report.docx"


TITLE = "DEVELOPMENT OF AN AI POWERED STUDENT ACADEMIC PERFORMANCE ANALYSIS AND ADVISORY SYSTEM"
AUTHOR = "DENNIS EMMANUEL"
MATRIC = "FOS/22/23/292155"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction, display=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_number(paragraph, roman=False):
    add_field(paragraph, "PAGE \\* roman" if roman else "PAGE", "1")


def set_page_number_start(section, start=1, fmt="decimal"):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))
    pg_num.set(qn("w:fmt"), fmt)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_table_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:keepNext")
    p_pr.append(node)


def avoid_split(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:keepLines")
    p_pr.append(node)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(10)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = style_name == "Heading 1"
        if style_name == "Heading 1":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for style_name in ("TOC 1", "TOC 2", "TOC 3"):
        if style_name in doc.styles:
            toc_style = doc.styles[style_name]
            toc_style.font.name = "Times New Roman"
            toc_style.font.size = Pt(12)
            toc_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            toc_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            toc_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            toc_style.paragraph_format.line_spacing = 1.0
            toc_style.paragraph_format.space_after = Pt(5)

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(11)
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.font.italic = False
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    if "Table Text" not in doc.styles:
        style = doc.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_text = doc.styles["Table Text"]
    table_text.font.name = "Times New Roman"
    table_text.font.size = Pt(10)
    table_text._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    table_text._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    table_text.paragraph_format.line_spacing = 1.0
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.first_line_indent = Inches(0)


def centred(doc, text="", size=12, bold=False, spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = spacing
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    return p


def body(doc, text, indent=True):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5) if indent else Inches(0)
    p.paragraph_format.space_after = Pt(4)
    avoid_split(p)
    return p


def heading(doc, text, level=2):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(14 if level > 1 else 12)
    p.paragraph_format.space_after = Pt(10)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(16)
            r.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.first_line_indent = Inches(0)
    spacer.paragraph_format.line_spacing = 1.0
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.keep_with_next = True
    # A nonbreaking space gives Word a real line box. A normal space may be
    # collapsed, which can allow the next paragraph to share the heading line.
    spacer_run = spacer.add_run("\u00a0")
    spacer_run.font.size = Pt(4)
    return p


def chapter(doc, number, title):
    p = heading(doc, f"CHAPTER {number} {title.upper()}", 1)
    p.paragraph_format.page_break_before = True


def numbered_list(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.hanging_indent = Inches(0.25)
        p.add_run(f"{index}. ").bold = True
        p.add_run(item)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.hanging_indent = Inches(0.25)
        p.add_run("• ")
        p.add_run(item)


def table(doc, number, title, headers, rows, widths=None, page_break_before=False):
    cap = doc.add_paragraph(f"Table {number}: {title}", style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.page_break_before = page_break_before
    keep_with_next(cap)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    prevent_table_row_split(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = value
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.style = doc.styles["Table Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = t.add_row().cells
        prevent_table_row_split(t.rows[-1])
        for i, value in enumerate(values):
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cells[i], "F2F6FA")
            p = cells[i].paragraphs[0]
            p.style = doc.styles["Table Text"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 20 else WD_ALIGN_PARAGRAPH.LEFT
            if widths:
                cells[i].width = Inches(widths[i])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.first_line_indent = Inches(0)
    spacer.paragraph_format.line_spacing = 1.0
    spacer.paragraph_format.space_after = Pt(6)
    return t


def figure(doc, number, path, caption, width=6.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("title", f"Figure {number}")
    picture._inline.docPr.set("descr", caption)
    c = doc.add_paragraph(f"Figure {number}: {caption}", style="Caption")
    c.paragraph_format.keep_with_next = False


def add_preliminaries(doc):
    section = doc.sections[0]
    configure_section(section)
    section.different_first_page_header_footer = True
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp, roman=True)
    set_page_number_start(section, 1, "lowerRoman")

    centred(doc, TITLE, 16, True, 1.35)
    for _ in range(2):
        centred(doc, "")
    centred(doc, "BY", 12, True)
    centred(doc, AUTHOR, 14, True)
    centred(doc, MATRIC, 12, True)
    for _ in range(2):
        centred(doc, "")
    centred(doc, "A PROJECT REPORT SUBMITTED TO THE DEPARTMENT OF COMPUTER SCIENCE", 12, True)
    centred(doc, "FACULTY OF SCIENCE", 12, True)
    centred(doc, "DELTA STATE UNIVERSITY ABRAKA", 12, True)
    centred(doc, "IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE AWARD OF THE DEGREE OF", 12, True)
    centred(doc, "BACHELOR OF SCIENCE IN COMPUTER SCIENCE", 12, True)
    for _ in range(2):
        centred(doc, "")
    centred(doc, "SEPTEMBER 2026", 12, True)

    heading(doc, "DECLARATION", 1)
    body(doc, "I declare that this project report was written by me and that it records the system I designed, implemented and evaluated. Sources used in the report are acknowledged in the text and listed in the references. The report has not been submitted in full for another degree or award.")
    signature = body(doc, "Name: Dennis Emmanuel\nMatriculation Number: FOS/22/23/292155\nSignature: ______________________________\nDate: __________________________________", indent=False)
    signature.alignment = WD_ALIGN_PARAGRAPH.LEFT
    signature.paragraph_format.line_spacing = 1.5

    heading(doc, "CERTIFICATION", 1)
    body(doc, "This project report titled Development of an AI Powered Student Academic Performance Analysis and Advisory System was carried out by Dennis Emmanuel, with matriculation number FOS/22/23/292155, in the Department of Computer Science, Delta State University, Abraka. The report has been examined and approved as meeting the requirements for the award of the Bachelor of Science degree in Computer Science.")
    table(doc, "P.1", "Certification signatures", ["Role", "Name", "Signature", "Date"], [
        ["Project supervisor", "", "", ""],
        ["Head of department", "", "", ""],
        ["External examiner", "", "", ""],
    ], [1.6, 2.0, 1.5, 1.1])

    heading(doc, "DEDICATION", 1)
    body(doc, "This work is dedicated to my family, teachers and classmates whose support made the project possible.")

    heading(doc, "ACKNOWLEDGEMENTS", 1)
    body(doc, "I am grateful to God for the strength to complete this project. I thank the lecturers in the Department of Computer Science, Delta State University, for the knowledge and guidance I received during my undergraduate study. I also appreciate my project supervisor for reviewing the work and directing the research. My thanks go to the course advisers who explained the structure of departmental broadsheets and the practical problems involved in sharing results. I am grateful to my family and friends for their patience and encouragement throughout the design, implementation and writing stages.")

    heading(doc, "ABSTRACT", 1)
    body(doc, "The Department of Computer Science at Delta State University uses result broadsheets to compile student grades, but a static file does not provide private access, automatic analysis or personal academic guidance. This project developed Compass, a web based student academic performance analysis and advisory system. The system uses a React interface, a FastAPI application server, Supabase Postgres and Auth, and a Groq hosted large language model. Course advisers can upload Excel broadsheets, review processed results and view class analytics. Students can view their own result history, GPA, CGPA, carryovers and notifications. The advisory agent retrieves academic records through defined backend tools before it produces an answer or runs a grade simulation. The project followed an iterative software development method. Testing covered the production build, code linting, authenticated browser flows, API responses and a read only inspection of the live database. The build completed successfully. Three measured runs produced mean response times of 35.4 ms for the health endpoint, 1.05 seconds for CGPA retrieval and 1.43 seconds for course retrieval. The adviser summary exceeded a 15 second test limit in three isolated runs, although it loaded in the browser. The live database contained 169 students, 1,293 results and 169 session baselines at the time of inspection. Security review found that some public table policies were too broad and that two privileged functions had mutable search paths. The system demonstrates that verified record retrieval can support useful academic guidance, while the test findings show that policy hardening and query optimisation are required before wider deployment.")

    heading(doc, "TABLE OF CONTENTS", 1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right click and update field to refresh the table of contents.")

    heading(doc, "LIST OF TABLES", 1)
    entries = [
        "Table 2.1  Comparison of related approaches", "Table 3.1  Functional requirements",
        "Table 3.2  Nonfunctional requirements", "Table 3.3  Core data dictionary",
        "Table 4.1  Development tools and versions", "Table 4.2  Build and lint results",
        "Table 4.3  API timing results", "Table 4.4  Functional test results",
        "Table 4.5  Live database evidence", "Table 4.6  Security and performance findings",
        "Table 5.1  Achievement of project objectives", "Table A.1  Selected FastAPI routes",
        "Table B.1  Five point grade mapping", "Table C.1  Test evidence summary",
    ]
    for item in entries:
        body(doc, item, indent=False)

    heading(doc, "LIST OF FIGURES", 1)
    entries = [
        "Figure 3.1  Use case model", "Figure 3.2  System architecture",
        "Figure 3.3  Result ingestion workflow", "Figure 3.4  Database relationship model",
        "Figure 3.5  Compass advisory sequence", "Figure 4.1  Public landing page",
        "Figure 4.2  Student dashboard", "Figure 4.3  Student academic record",
        "Figure 4.4  Compass start screen", "Figure 4.5  Verified Compass response",
        "Figure 4.6  Adviser analytics dashboard", "Figure 4.7  Adviser upload form",
        "Figure 4.8  Upload history",
    ]
    for item in entries:
        body(doc, item, indent=False)

    heading(doc, "LIST OF ABBREVIATIONS", 1)
    table(doc, "P.2", "Abbreviations used in the report", ["Abbreviation", "Meaning"], [
        ["AI", "Artificial intelligence"], ["API", "Application programming interface"],
        ["CGPA", "Cumulative grade point average"], ["DELSU", "Delta State University"],
        ["FK", "Foreign key"], ["GPA", "Grade point average"],
        ["JWT", "JSON Web Token"], ["LLM", "Large language model"],
        ["MIS", "Management information system"], ["PK", "Primary key"],
        ["RBAC", "Role based access control"], ["REST", "Representational State Transfer"],
        ["RLS", "Row level security"], ["UAT", "User acceptance testing"],
    ], [1.6, 4.6])


def chapter_one(doc):
    chapter(doc, "ONE", "INTRODUCTION")
    heading(doc, "1.1 Background to the study", 2)
    body(doc, "Universities depend on accurate academic records for progression decisions, graduation checks and student support. Result management has therefore moved from paper registers toward database systems that can store grades and produce reports. A student portal adds a personal access layer to this process. It allows each student to review performance without searching through a class document. Learning analytics extends this idea by turning stored results into summaries that can support early intervention (Namoun & Alshanqiti, 2021; Albreiki et al., 2021).")
    body(doc, "The value of a portal depends on availability, privacy and clarity. A system may compute results correctly and still fail students when the access channel is unreliable. In many institutions, departments use spreadsheets because the format is familiar and supports manual verification. When the official portal is unavailable, advisers may export the spreadsheet as a PDF and share it through a class messaging group. That method is fast, but every recipient can see the records of other students. A large static file is also difficult to search on a phone and cannot explain how a GPA was obtained.")
    body(doc, "Recent artificial intelligence research offers another way to present academic records. Transformer models process relationships between words through attention mechanisms (Vaswani et al., 2017). Large language models built on this foundation can follow instructions and produce useful text from a short prompt (Brown et al., 2020; Zhao et al., 2023). Their output is not automatically reliable. A model can state an unsupported figure because its response is generated from learned patterns. Studies of hallucination describe this as a continuing risk, especially when a factual answer must match an external record (Ji et al., 2023).")
    body(doc, "Tool use reduces this risk by separating record retrieval from language generation. ReAct combines reasoning steps with actions in an external environment (Yao et al., 2023). Toolformer and later surveys show how a model can call a calculator, database or other service when a task needs information beyond the prompt (Schick et al., 2023; Qin et al., 2023). In an academic advisory system, a tool can retrieve a student's CGPA and courses from the database. The language model can then explain those values. This method improves grounding, although access control, tool correctness and clear limits remain necessary.")
    body(doc, "This project applies that approach to the Department of Computer Science at Delta State University, Abraka. Compass accepts the Excel format already used by course advisers, processes the rows and stores the resulting records in Supabase. Students receive an individual dashboard and can ask questions about their verified records. Advisers receive upload, history and class analytics interfaces. The system is a departmental prototype, not a replacement for the university's official senate approved result process.")

    heading(doc, "1.2 Statement of the problem", 2)
    body(doc, "Students in the department need a dependable and private way to view their results. The existing fallback is a broad class file shared through a group channel. It exposes grades and CGPA values to people who only need their own record. It also forces a student to locate a row manually and interpret the result without support. The file does not provide automatic alerts, carryover tracking or a safe way to examine the effect of possible future grades.")
    body(doc, "Course advisers face a related problem. A broadsheet contains many scores, units and calculated totals. Manual formulas can be copied incorrectly, and a small error can affect a semester GPA or CGPA. The current sharing method gives no structured preview that reports missing values, invalid scores or a difference between uploaded and recomputed totals. The department therefore needs a system that protects individual access, checks result data and provides explanations based on the stored record.")

    heading(doc, "1.3 Aim and objectives", 2)
    body(doc, "The aim of this project is to develop an AI powered student academic performance analysis and advisory system for the Department of Computer Science at Delta State University.")
    body(doc, "The objectives are to:", indent=False)
    numbered_list(doc, [
        "design a role based web platform that gives students personal result access and gives course advisers an administrative workspace;",
        "develop a result ingestion and anomaly checking process for departmental Excel broadsheets;",
        "implement an AI advisory agent that retrieves stored academic data and supports grade simulations;",
        "evaluate the implemented system through build checks, functional inspection, API tests and database review.",
    ])

    heading(doc, "1.4 Research questions", 2)
    numbered_list(doc, [
        "How can departmental broadsheets be converted into private student records without changing the adviser's familiar Excel workflow?",
        "How can the system detect invalid values and differences in calculated academic totals before an upload is confirmed?",
        "How can an AI adviser answer student questions from verified records instead of relying only on generated text?",
        "What do functional, performance and security checks show about the readiness of the prototype?",
    ])

    heading(doc, "1.5 Significance of the study", 2)
    body(doc, "Students gain one place to view GPA, CGPA, course scores, degree classification and outstanding courses. The layout is designed for a phone sized screen as well as a desktop browser. Notifications inform a student when new records are available. Compass also explains a stored result in ordinary language and can calculate a possible outcome from values supplied by the student.")
    body(doc, "Course advisers gain a controlled upload workflow and a class level summary. The preview stage separates file selection from final confirmation, which gives the adviser time to inspect processed rows. The analytics interface shows the number of students, average CGPA, degree class distribution, students below a selected threshold and carryover counts. These summaries can support follow up, but they do not replace academic judgement.")
    body(doc, "The project also provides a practical case study of agent tool use in a Nigerian university setting. It connects current AI research to an operational information system and records the limitations found during testing. This is important because a persuasive interface does not prove that a system is secure or ready for institution wide use.")

    heading(doc, "1.6 Scope of the study", 2)
    body(doc, "The system covers the Department of Computer Science at Delta State University. It processes Microsoft Excel broadsheets in wide and long layouts, computes grades on a five point scale, stores course results, calculates GPA and CGPA, tracks failed courses and produces class summaries. The user roles represented in the working application are student, course adviser and administrator. The screenshots in this report focus on the student and adviser roles because valid test accounts were available for those roles.")
    body(doc, "Compass is limited to academic result questions and calculations supported by backend tools. It does not approve results, change a grade, register courses, pay fees or make a final graduation decision. The evaluation uses one authenticated student account, one authenticated adviser account, a live database inspection and developer run tests. It does not claim a controlled usability experiment with a representative student sample.")

    heading(doc, "1.7 Limitations of the study", 2)
    body(doc, "The live data inspected for this report represents one session and one semester. The test environment also depends on external Supabase and Groq services, so network conditions affect response time. The current adviser dashboard performs several database operations and exceeded the isolated 15 second timing limit in three runs. The database review found policies and privileged functions that require hardening before a public deployment. These limitations are included in the findings and recommendations.")

    heading(doc, "1.8 Definition of terms", 2)
    terms = [
        ("Academic advisory agent", "A software agent that retrieves student records and produces academic explanations or calculations within defined limits."),
        ("Broadsheet", "An Excel worksheet that contains the results of many students and the calculations used by a course adviser."),
        ("Carryover", "A failed course that a student must retake under the applicable academic rules."),
        ("CGPA", "The cumulative average of grade points earned over the completed academic periods represented in the system."),
        ("GPA", "The weighted average of grade points for a defined semester."),
        ("Grounding", "The use of retrieved evidence to support a generated answer."),
        ("RLS", "A Postgres control that applies row access rules to database queries."),
        ("Tool calling", "A process in which a language model requests a defined software function and uses the returned result in its answer."),
    ]
    for term, meaning in terms:
        p = body(doc, "", indent=False)
        p.add_run(term + ": ").bold = True
        p.add_run(meaning)


def chapter_two(doc):
    chapter(doc, "TWO", "LITERATURE REVIEW")
    body(doc, "This chapter reviews the concepts, theories and related studies that support the design of Compass. The review covers result management, learning analytics, large language models, tool using agents, privacy and access control. It then compares related approaches and states the gap addressed by the project.")

    heading(doc, "2.1 Conceptual review", 2)
    heading(doc, "2.1.1 Academic result management systems", 3)
    body(doc, "A result management system receives assessment data, applies academic rules and stores the outcome for later use. Its basic functions are data entry, validation, calculation, retrieval and reporting. A spreadsheet can perform several of these tasks, but it does not provide identity based access when the whole file is distributed. A web system can separate the stored record from the way each user sees it. This separation supports personal views and makes it easier to correct presentation errors without altering the source grades.")
    body(doc, "The quality of an academic system depends on more than a correct formula. ISO/IEC 25010:2023 groups software quality around characteristics that include functional suitability, performance efficiency, usability, reliability and security. Those ideas are relevant to result systems because a calculation may be accurate while access is slow, confusing or unsafe. The evaluation in this project therefore includes behaviour, timing and database security evidence.")

    heading(doc, "2.1.2 Learning analytics and student performance", 3)
    body(doc, "Learning analytics uses data about learners to understand and improve learning. Systematic reviews show that student performance prediction often uses demographic data, prior grades, attendance and activity records (Namoun & Alshanqiti, 2021; Albreiki et al., 2021). The common goal is early identification of students who may need support. Kabathova and Drlik (2021), for example, compared machine learning methods for university dropout prediction. Such work shows the value of patterns in academic data, but a prediction is not the same as a verified result.")
    body(doc, "Compass does not train a predictive model on the department's records. Its class analytics uses direct calculations, thresholds and distributions. This choice keeps the result explainable. A student marked below a CGPA threshold can be traced to stored values instead of an opaque classification. The grade simulator is also deterministic. It calculates the effect of proposed scores without claiming that the scores will occur.")

    heading(doc, "2.1.3 Large language models", 3)
    body(doc, "The Transformer introduced an attention based architecture that can process relationships across a sequence without relying on recurrent steps (Vaswani et al., 2017). BERT showed how pretraining could produce useful language representations for later tasks (Devlin et al., 2019). GPT style models demonstrated strong few shot behaviour at a much larger scale (Brown et al., 2020). Libraries such as Transformers made pretrained models easier to study and deploy (Wolf et al., 2020).")
    body(doc, "An LLM predicts likely tokens. It does not contain a guaranteed copy of the current university database. Bender et al. (2021) warned that fluent language can encourage readers to assign understanding or authority that the model does not possess. Ji et al. (2023) reviewed hallucination in natural language generation and described the gap between a plausible output and its source evidence. These concerns matter in academic advising because an incorrect CGPA can affect a student's decisions.")

    heading(doc, "2.1.4 Tool using and agentic AI", 3)
    body(doc, "Chain of thought prompting can improve performance on some reasoning tasks by encouraging intermediate steps (Wei et al., 2022). Intermediate reasoning alone cannot supply a missing institutional fact. ReAct joins reasoning with actions, allowing a model to obtain observations from an environment before continuing (Yao et al., 2023). Toolformer explored how models can learn to call external tools (Schick et al., 2023). Qin et al. (2023) organised tool learning into task planning, tool selection, tool use and response generation.")
    body(doc, "AgentBench evaluates language models in environments where they must act rather than answer a single static question (Liu et al., 2024). Its findings support the need to test an agent as a complete process. A correct model response depends on the tool definition, retrieved data, access checks and final wording. Compass follows this process by exposing academic functions to the agent. The backend remains responsible for calculations and database queries.")

    heading(doc, "2.1.5 Conversational support in education", 3)
    body(doc, "The release of general purpose conversational models led to rapid study in education. Lo (2023) found both possible benefits and concerns in early research on ChatGPT. Tlili et al. (2023) examined educational uses and ethical issues through a case based analysis. Kasneci et al. (2023) discussed opportunities for personalised explanation as well as risks involving bias, incorrect information and assessment practice. UNESCO (2023) recommends human oversight, data protection and age appropriate use.")
    body(doc, "Educational support can be helpful when the system states its purpose and keeps important decisions with qualified staff. Compass provides explanations and calculations. It does not act as a lecturer, counsellor or senate authority. The interface labels it as an academic adviser, while the implementation limits its accessible tools to the student's record and academic simulations.")

    heading(doc, "2.1.6 Privacy and access control", 3)
    body(doc, "Academic results are personal records. Sharing a class broadsheet gives each recipient more information than required for their own purpose. The Nigeria Data Protection Act 2023 establishes duties concerning lawful processing, security and the rights of data subjects. A departmental system should therefore minimise exposure and keep access tied to a valid identity.")
    body(doc, "Role based access control assigns permissions through roles such as student or adviser (Sandhu et al., 1996; Ferraiolo et al., 2001). Authentication confirms an identity, while authorisation decides what that identity may do. Supabase Auth provides accounts and sessions. Postgres row level security can add database rules to each table. Supabase documentation states that RLS must be enabled for tables in an exposed schema and that ownership checks should use the authenticated user identifier. The present project uses role aware interfaces and backend checks, but the live policy audit shows that several database policies remain too broad.")

    heading(doc, "2.1.7 Explainability in educational systems", 3)
    body(doc, "An educational system should allow a user to understand how an important value was obtained. Explainability is especially relevant when a result is used to advise a student. Khosravi et al. (2022) describe explainable artificial intelligence in education as a field that connects technical explanation with the needs of learners and educators. A useful explanation should therefore identify the relevant evidence and present it in language that fits the user's task. A long technical description may be accurate but still fail to help a student decide what to do next.")
    body(doc, "Compass uses two forms of explanation. The first is deterministic. GPA and CGPA values can be traced to units, grade points and cumulative baselines. The second is conversational. The advisory agent places retrieved values in a short response and can explain outstanding courses or a simulated result. This design does not expose hidden model reasoning. It exposes the academic inputs and the calculation path that matter to the user. The approach is consistent with the view that explanations should support checking and informed action instead of serving as decoration.")

    heading(doc, "2.1.8 Responsible use of generative AI", 3)
    body(doc, "Generative AI creates new risks when it is connected to personal records. A response may be fluent even when its factual basis is weak. It may also reveal data that the current user should not receive. NIST (2023) treats validity, reliability, privacy, transparency and accountability as connected parts of AI risk management. UNESCO (2023) also recommends human oversight and protection of learner data. These principles are relevant to a student adviser because an incorrect statement may influence course planning or a student's view of academic standing.")
    body(doc, "The responsible design boundary for Compass is narrow. The agent may explain records and perform defined simulations, but it may not change a grade, approve graduation or replace a qualified course adviser. Academic calculations remain in tested functions. The student supplies the question, while the system retrieves only the information needed for the response. The interface and report also state the limits of the prototype. This boundary supports useful interaction while keeping final academic authority with the department.")

    heading(doc, "2.2 Theoretical framework", 2)
    heading(doc, "2.2.1 ReAct framework", 3)
    body(doc, "ReAct is the main theory for the advisory process. A student question starts a cycle in which the model identifies the information needed, selects an available tool, receives an observation and prepares the answer (Yao et al., 2023). In Compass, a question about current standing calls the full academic record or GPA tool. A simulation calls a deterministic calculation tool. The model cannot directly write to the result tables.")
    body(doc, "The framework is suitable because it makes the source of a factual answer visible in the system design. Its use does not make every answer correct. An incorrect query, weak access rule or unclear user question can still produce an error. The implementation therefore combines tool use with scope instructions and server side functions.")

    heading(doc, "2.2.2 Role based access control model", 3)
    body(doc, "The second framework is RBAC. Students need read access to their own result and chat history. Advisers need broader access to the level and department assigned to them, together with upload and analysis actions. Administrators manage adviser approval. These permissions form distinct job roles. The model is useful for interface routing and API decisions. Database RLS must then express the same ownership rules to prevent a user from bypassing the interface.")

    heading(doc, "2.2.3 Client server and REST architecture", 3)
    body(doc, "Compass separates the React client from the FastAPI server. Fielding (2000) described REST as an architectural style built around resources, stateless interaction and a uniform interface. The client sends HTTP requests for profiles, results, uploads, analytics and agent messages. The server validates input and returns JSON. This separation allows the interface and business logic to be tested independently.")

    heading(doc, "2.3 Review of related works", 2)
    body(doc, "Namoun and Alshanqiti (2021) reviewed data mining and learning analytics studies concerned with student performance. They found frequent use of classification methods and academic features. The review provides a broad account of prediction practice, although it does not describe a student controlled advisory agent grounded in a university result store.")
    body(doc, "Albreiki et al. (2021) reviewed machine learning methods for student performance prediction. Their work records the algorithms, datasets and evaluation measures used across studies. It supports the need for careful evaluation and explains why a class prediction should not be presented without a method and test set. Compass therefore avoids calling its threshold summaries predictions.")
    body(doc, "Kabathova and Drlik (2021) studied dropout prediction in university courses with several machine learning methods. The study focused on identifying risk from course data. The output is useful to administrators, but it does not solve private result delivery or allow a student to ask a record specific question.")
    body(doc, "Zawacki-Richter et al. (2019) reviewed artificial intelligence applications in higher education and observed that much of the published work focused on profiling, prediction, assessment and adaptive systems. They also noted limited educator involvement in parts of the research. Their review helps place Compass between administrative analytics and direct student support.")
    body(doc, "Ouyang et al. (2022) reviewed empirical work on artificial intelligence in online higher education. Their analysis shows that AI systems are used for prediction, recommendation and automated support. Compass differs by using a language model as the explanation layer while keeping academic values in a relational database and deterministic functions.")
    body(doc, "Park et al. (2023) created generative agents that store observations, reflect and plan in a simulated environment. The study demonstrates how memory can support coherent interaction. Compass uses a narrower form of memory through chat sessions and messages. It does not attempt to model a student's personality or behaviour.")
    body(doc, "Kasneci et al. (2023) examined the opportunities and challenges of large language models in education. They identified possible value in personalised support and explanation, together with risks involving bias, inaccurate output and inappropriate reliance. Their analysis supports a cautious design in which a language model assists a defined task and a human institution remains responsible for decisions. Compass applies this principle by limiting the agent to result explanation and simulations based on stored data.")
    body(doc, "Khosravi et al. (2022) reviewed explainable artificial intelligence in education and showed that explanation needs differ across students, instructors and administrators. Their work is relevant to the way Compass separates a student record page from the adviser analytics page. Students receive personal values and plain explanations, while advisers receive class summaries that support follow up. The current prototype does not yet evaluate whether the explanations improve understanding, so this remains a subject for a controlled user study.")
    body(doc, "Liu et al. (2024) introduced AgentBench to test language models in environments that require actions and observations. The study shows that an agent should be evaluated as a sequence rather than as a single text response. This idea informed the inspection of the Compass response path. The test checked login, session retrieval, database values and the displayed answer. The scope was small, but it examined the connected process that produces the advice.")
    body(doc, "Schick et al. (2023) and Qin et al. (2023) show how language models can use external tools. Tool access gives a model current observations, but the result still depends on the tool definition and the permission applied to it. Compass therefore treats each academic tool as part of the application layer. The model does not receive a general database connection. This distinction is important because safe tool design requires both correct computation and restricted access.")

    table(doc, "2.1", "Comparison of related approaches", ["Approach", "Personal access", "Verified record retrieval", "Interactive advice", "Upload checking", "Main limitation"], [
        ["Shared PDF broadsheet", "No", "Static file", "No", "No", "Exposes class records and requires manual search"],
        ["Standard student portal", "Yes", "Yes", "Limited", "Varies", "Usually presents a read only table"],
        ["Performance prediction study", "Usually administrative", "Dataset based", "No", "No", "Prediction may not explain the official record"],
        ["General AI tutor", "Yes", "No", "Yes", "No", "May produce unsupported institutional facts"],
        ["Compass", "Yes", "Yes through tools", "Yes", "Yes", "Prototype still needs policy hardening and optimisation"],
    ], [1.15, 0.8, 1.1, 0.85, 0.85, 1.7], page_break_before=True)

    heading(doc, "2.4 Research gap", 2)
    body(doc, "The literature shows strong interest in predicting student outcomes and using conversational AI for educational support. These lines of work often remain separate. Prediction systems usually serve administrators, while general chat systems do not have controlled access to the institution's result database. Standard portals protect personal access but give little explanation or simulation support.")
    body(doc, "The local workflow adds a practical gap. Course advisers already maintain structured Excel broadsheets, yet students receive a class wide static document when a personal portal is unavailable. Compass addresses the gap by joining Excel ingestion, private result presentation, class analytics and a tool using advisory agent in one departmental prototype. The contribution is the integration and its tested implementation, not a new language model or prediction algorithm.")

    heading(doc, "2.5 Summary of the literature review", 2)
    body(doc, "The review established the need for accurate result management, limited disclosure and evidence based advisory responses. Learning analytics explains the use of summaries and risk thresholds. Transformer and agent research explains the conversational layer and the retrieval cycle. RBAC and RLS explain the controls required for academic data. Related works leave room for an integrated departmental system that preserves the existing broadsheet workflow while giving each student an interactive personal view.")


def chapter_three(doc):
    chapter(doc, "THREE", "SYSTEM ANALYSIS AND DESIGN")
    heading(doc, "3.1 Development methodology", 2)
    body(doc, "The project used an iterative development method based on Agile principles (Beck et al., 2001). Work was divided into small functional parts: database and authentication, result ingestion, student views, adviser analytics, notifications and the advisory agent. Each part was tested before the next part was joined to it. This suited the project because the real broadsheet format revealed requirements that were difficult to define at the start.")
    body(doc, "The method also supported correction. Result files could be parsed early, then the GPA and carryover rules were adjusted after comparison with the source workbook. Interface pages were tested with authenticated accounts while backend routes remained visible in their own logs. The documentation stage reviewed the final code and live database so the report describes the current system instead of an early design.")
    body(doc, "The work followed five practical cycles. The first cycle established the Supabase project, application profiles and authentication flow. The second cycle processed Excel data and compared calculated grades with the workbook. The third cycle created student and adviser pages. The fourth cycle connected the Compass agent to academic tools and chat history. The final cycle reviewed build output, browser behaviour, API timing and database configuration. A problem found in one cycle was corrected before later features depended on it.")
    body(doc, "Requirements were checked against observable evidence. A page was considered implemented when the correct authenticated role could reach it and its main data loaded. A calculation was checked by comparing the displayed value with stored course values or a database summary. Database security was evaluated separately from interface behaviour because a hidden page does not prevent direct access to a public data endpoint. This separation made it possible to report that the user flows worked while some database controls still required correction.")

    heading(doc, "3.2 Analysis of the existing system", 2)
    body(doc, "The existing departmental process starts when a course adviser receives scores and enters them into an Excel broadsheet. The workbook calculates grades and summary values. The adviser checks the file and distributes a PDF copy through a student group when the main result portal is unavailable. Each student opens the same file and searches for a matriculation number.")
    body(doc, "The process has useful qualities. Excel is familiar, works offline and supports visible formulas. It also has weaknesses. Distribution is class wide, search is manual, corrections require another file, and a static document cannot notify only affected students. There is no conversational support or structured history of a student's questions.")
    body(doc, "The existing process also joins several responsibilities in one file. The spreadsheet acts as an input form, calculator, store and publication document. This makes the file easy to begin with, but it becomes difficult to control after distribution. A student may keep an old copy after a correction. A formula can be overwritten without an access log. The adviser must also answer repeated questions about CGPA, failed courses and degree classification even when the required values are already present in the sheet.")
    body(doc, "The main information flow can be described in six steps. Scores are collected, entered into course columns, converted to grades, combined with units, checked by the adviser and exported for students. Privacy is weakest at the last step because the exported file contains the class. Reliability risks appear earlier when column names, scores, units or formulas are inconsistent. The proposed system must address both points. It must validate the imported structure and then give each authenticated student a limited view.")

    heading(doc, "3.3 Analysis of the proposed system", 2)
    body(doc, "The proposed system keeps Excel as the adviser input but changes the delivery path. The file is uploaded through a form that records semester and session. The backend parses rows, normalises course fields, validates scores and units, recomputes grade values and prepares a preview. Confirmation writes the upload and result records. Students then see only the record linked to their account. Notifications and Compass use the same stored data.")
    body(doc, "The proposed flow separates preparation from publication. Selecting a file does not immediately change the database. The preview route reads the workbook and returns processed rows, summary values and validation messages. The adviser can stop when the file or selected period is wrong. Confirmation is a separate request that persists the approved rows. This arrangement reduces accidental publication and creates a clear point at which an upload record can be created.")
    body(doc, "The web system also separates presentation from academic logic. React components display values but do not define the grade scale. FastAPI services validate the workbook and calculate results. Supabase stores the confirmed state and provides the account identity. The Groq model receives selected observations through agent tools. Each layer has a limited responsibility, which makes faults easier to locate and gives the database and calculations a life outside the chat interface.")
    body(doc, "Figure 3.1 shows the main use cases available to the student and course adviser.")
    figure(doc, "3.1", DIAGRAM / "use_case.png", "Use case model for student and adviser roles")

    heading(doc, "3.4 Requirements specification", 2)
    heading(doc, "3.4.1 Functional requirements", 3)
    table(doc, "3.1", "Functional requirements", ["ID", "Requirement", "Primary role", "Acceptance condition"], [
        ["FR1", "Authenticate a user with email and password", "All roles", "A valid account reaches its authorised portal"],
        ["FR2", "Display a student's academic record", "Student", "The page shows stored CGPA, courses and outstanding items"],
        ["FR3", "Upload and preview an Excel broadsheet", "Adviser", "The system parses supported layouts and reports invalid rows"],
        ["FR4", "Confirm processed result data", "Adviser", "Upload and result records are stored after confirmation"],
        ["FR5", "Calculate GPA, CGPA and carryovers", "System", "Values follow configured units, scores and grade points"],
        ["FR6", "Show class analytics", "Adviser", "Dashboard presents counts, averages, distributions and risk lists"],
        ["FR7", "Create and read notifications", "System and student", "Affected students can view result notices"],
        ["FR8", "Answer record based questions", "Student", "Compass retrieves data through defined tools"],
        ["FR9", "Store chat sessions and messages", "Student", "A signed in student can return to saved conversations"],
        ["FR10", "Approve or revoke adviser access", "Administrator", "Adviser status controls entry to the adviser portal"],
    ], [0.45, 2.2, 1.0, 2.65])

    heading(doc, "3.4.2 Nonfunctional requirements", 3)
    table(doc, "3.2", "Nonfunctional requirements", ["Area", "Requirement", "Design response", "Measure"], [
        ["Security", "Protect personal academic data", "Authentication, role routing, backend checks and RLS", "Policy audit and role based tests"],
        ["Usability", "Support phone and desktop access", "Responsive layouts with short labels and clear actions", "Browser inspection at 1280 by 720 and responsive CSS"],
        ["Performance", "Return ordinary student records within a few seconds", "REST endpoints and indexed identifiers", "Three run response timing"],
        ["Reliability", "Reject malformed input before persistence", "Preview and validation stage", "Valid and invalid file tests"],
        ["Maintainability", "Separate interface, routes and services", "React pages, FastAPI routers and service modules", "Code structure review"],
        ["Compatibility", "Accept current departmental workbooks", "Wide and long layout parsing", "Workbook ingestion checks"],
    ], [1.0, 1.8, 2.25, 1.25])

    heading(doc, "3.5 System architecture", 2)
    body(doc, "Compass uses a three layer web architecture with external managed services. The React client handles navigation and display. FastAPI exposes routes for authentication profiles, students, uploads, results, analytics, notifications and the advisory agent. Supabase provides Postgres storage, authentication and realtime events. The Groq API hosts the language model used for response generation. Figure 3.2 presents these parts and their main data movement.")
    figure(doc, "3.2", DIAGRAM / "architecture.png", "Logical architecture of the implemented system")
    body(doc, "The browser holds the Supabase session and sends the authenticated user identifier to relevant backend routes. The backend uses service functions to read and process records. This design places academic calculations outside the language model. The database is the source for stored results. The model receives selected data through tools when it needs to explain a record.")

    heading(doc, "3.6 Result ingestion design", 2)
    body(doc, "The ingestion service accepts an Excel file together with semester and session. It recognises a wide broadsheet, where course values appear across columns, and a long layout, where each row represents a course result. Normalisation converts both forms to common result fields. Validation checks required identifiers, score ranges, units and course codes. The service then derives grades, grade points and quality points.")
    body(doc, "The preview step is important because an uploaded file may contain a structural or calculation error. The adviser can inspect the parsed result before confirmation. On confirmation, the service stores the upload, courses, students and result rows, updates session baselines and creates notifications. Figure 3.3 shows the workflow.")
    figure(doc, "3.3", DIAGRAM / "ingestion_workflow.png", "Result ingestion and validation workflow")

    heading(doc, "3.7 GPA and CGPA calculation", 2)
    body(doc, "For a course with unit value U and grade point G, the quality point is U multiplied by G. Semester GPA is the sum of quality points divided by the sum of attempted units for that semester. CGPA is the cumulative quality point total divided by the cumulative unit total represented by the current results and stored baseline. The five point mapping used by the system is A equals 5, B equals 4, C equals 3, D equals 2 and F equals 0.")
    body(doc, "Session baselines preserve cumulative totals that existed before the current upload. This avoids treating one imported semester as a complete degree record. The live database contained one baseline row for each of 169 students, and a consistency query found that the current master baseline values matched the session baseline values at the time of review.")
    body(doc, "A simple example shows the calculation. Assume a student earns A in a three unit course, B in a two unit course and C in a three unit course. The quality points are 15, 8 and 9. The semester total is 32 quality points across eight units, so the GPA is 4.00. The application performs the same weighted calculation for every valid course row. It does not average the letter grades or the raw scores directly.")
    body(doc, "For cumulative calculation, the current semester totals are added to the previous baseline totals. If the previous record contains 120 quality points across 30 units and the new semester contributes 32 quality points across eight units, the cumulative value is 152 divided by 38, which is 4.00. Keeping units and quality points is more reliable than averaging two GPA values because semesters may contain different unit totals.")
    body(doc, "Carryovers are derived from failed compulsory courses. An F grade is mapped to zero grade points and the course remains outstanding until a later result satisfies the rule used by the system. The interface distinguishes previous outstanding courses from current results. This allows a student to see that a past failure still requires attention even when the current semester GPA is high.")

    heading(doc, "3.8 Database design", 2)
    body(doc, "The database uses UUID primary keys for operational entities. Students and advisers store the Supabase Auth user identifier that links an account to an application profile. Results link students, courses and uploads. Notifications link to students. Chat messages link to chat sessions. Figure 3.4 presents the main relationships observed in the live Supabase schema.")
    figure(doc, "3.4", DIAGRAM / "erd.png", "Main database entities and relationships")
    table(doc, "3.3", "Core data dictionary", ["Table", "Purpose", "Important fields", "Main relationship"], [
        ["students", "Student identity and cumulative baseline", "id, auth_user_id, matric_number, name, level", "One student has many results and notifications"],
        ["advisers", "Approved adviser profile", "id, auth_user_id, department, level, status", "One adviser has many uploads"],
        ["courses", "Course catalogue", "id, course_code, title, units, course_type", "One course has many results"],
        ["results", "Individual course outcome", "student_id, course_id, score, grade, session", "Links a student, course and upload"],
        ["uploads", "Broadsheet import record", "adviser_id, filename, semester, session", "Parent of imported results"],
        ["notifications", "Student alert", "student_id, message, is_read", "Many notices belong to a student"],
        ["chat_sessions", "Conversation header", "matric_number, title, is_temporary", "One session has many messages"],
        ["chat_messages", "Conversation turn", "session_id, role, content", "Belongs to one chat session"],
        ["student_session_baselines", "Cumulative total for a session", "student_id, session, units, points", "Many baselines can belong to a student"],
    ], [1.1, 1.65, 2.4, 1.45])

    heading(doc, "3.9 Advisory agent design", 2)
    body(doc, "The advisory route receives the student's matriculation number, the current message and optional chat history. The system prompt defines the academic scope and the available tools. Tools retrieve the full record, course performance, carryovers, semester GPA and cumulative GPA. Other tools calculate future GPA and the effect of proposed grades. The agent can make more than one tool call before it responds.")
    body(doc, "The sequence in Figure 3.5 shows that the generated answer follows record retrieval. The design improves traceability because a numerical statement can be checked against the tool result. It also restricts the agent from performing a result update.")
    figure(doc, "3.5", DIAGRAM / "agent_sequence.png", "Sequence used to produce a grounded Compass response")

    heading(doc, "3.10 Security design", 2)
    body(doc, "Supabase Auth handles password authentication and session tokens. The client detects whether the account has a student, adviser or administrator profile and routes it to the matching interface. Adviser routes check status and assigned level. Student result routes use the matriculation number linked to the profile. Secrets used by the server are stored in environment files and are not included in the frontend bundle.")
    body(doc, "Database RLS is intended to provide another control below the API. During final inspection, the chat policies were scoped through auth.uid(), but several main application tables had permissive policies with conditions that evaluate to true. The design requirement is therefore only partly met in the current database. Chapter Four records the finding, and Chapter Five recommends ownership based policies and removal of unnecessary public grants.")
    body(doc, "The main security subjects are authentication, authorisation, input validation, secret handling and auditability. Authentication reduces anonymous access, but it does not decide whether a student may read another student's row. That decision belongs to authorisation rules in the API and database. Input validation protects result integrity by rejecting missing identifiers and values outside permitted ranges. Secret handling keeps service credentials outside source files delivered to the browser.")
    body(doc, "The design assumes that a browser client can be modified by its user. A menu item hidden by React is therefore a usability control, not a complete security control. Server routes must derive or verify the student's identity, and database policies must reject rows outside that identity. Adviser access also requires an active profile and an assigned departmental scope. Administrative actions require a separate role. These checks should produce the same decision at every layer.")
    body(doc, "Auditability is provided partly through upload records, chat sessions and notifications. An upload record identifies the adviser, file name, period and result rows created from the file. Chat messages retain the question and response for a signed in student. These records support investigation, but the prototype does not yet provide a complete immutable audit log for policy changes, deletions and administrator actions. That feature is recommended before wider deployment.")

    heading(doc, "3.11 Interface design", 2)
    body(doc, "The student interface uses a narrow centred layout with four main destinations: Home, Results, Compass and Settings. The adviser interface uses a sidebar because it is designed for wider administrative tasks. Blue is used for primary actions and selected navigation. Cards separate summaries without placing long explanations inside the interface. These decisions support recognition and consistency, two principles discussed in usability engineering (Nielsen, 1994).")

    heading(doc, "3.12 Chapter summary", 2)
    body(doc, "This chapter analysed the existing broadsheet process and translated its needs into a layered web system. It defined functional and quality requirements, described the ingestion and calculation rules, modelled the database and explained the advisory sequence. The next chapter presents the implemented interfaces and the evidence obtained during testing.")


def chapter_four(doc):
    chapter(doc, "FOUR", "IMPLEMENTATION TESTING AND RESULTS")
    heading(doc, "4.1 Implementation environment", 2)
    body(doc, "The system was implemented as a JavaScript client and Python server connected to managed cloud services. Table 4.1 lists the versions found in the project configuration during the final review. Package lock files are retained in the repository so that dependency versions can be reproduced.")
    table(doc, "4.1", "Development tools and versions", ["Component", "Technology", "Observed version or role"], [
        ["Frontend", "React", "19.2.7"], ["Build tool", "Vite", "8.1.2 during build"],
        ["Styling", "Tailwind CSS", "4.3.2"], ["Routing", "React Router", "7.18.1"],
        ["Backend", "FastAPI", "Python application framework"], ["Server", "Uvicorn", "Local ASGI server"],
        ["Database and auth", "Supabase", "Hosted Postgres 17 and email authentication"],
        ["Data processing", "pandas and openpyxl", "Excel parsing and worksheet access"],
        ["AI provider", "Groq", "Hosted inference API"], ["AI model", "openai/gpt-oss-120b", "Configured advisory model"],
    ], [1.55, 2.0, 2.75])

    heading(doc, "4.2 Frontend implementation", 2)
    body(doc, "The public page introduces the purpose of Compass and provides links to sign up or sign in. Figure 4.1 shows the landing page served by the local Vite development server.")
    figure(doc, "4.1", SCREEN / "01_landing_page.png", "Public landing page", 6.2)
    body(doc, "After authentication, the student home page presents three tasks: view the academic record, ask Compass and read notifications. The bottom navigation remains available on the main student pages. Figure 4.2 shows the authenticated dashboard for the test student account.")
    figure(doc, "4.2", SCREEN / "02_student_dashboard.png", "Authenticated student dashboard", 6.2)
    body(doc, "The academic record page retrieves cumulative and course data from separate FastAPI endpoints. It displays CGPA, degree class, total courses, total units, sessions and outstanding courses. Course results are grouped by session and semester. Figure 4.3 shows the test account with a CGPA of 4.03 and eight current semester courses.")
    figure(doc, "4.3", SCREEN / "03_student_results.png", "Student academic record page", 6.2)
    body(doc, "Compass opens as a dedicated chat interface with saved sessions, temporary chat and a message input. Figure 4.4 shows the start screen. The student can begin a new question or return to an earlier conversation.")
    figure(doc, "4.4", SCREEN / "04_compass_start.png", "Compass advisory start screen", 6.2)
    body(doc, "Figure 4.5 shows an existing response to a question about CGPA and carryovers. The displayed CGPA of 4.03 and the two previous outstanding courses matched the academic record page during inspection. This is a single verified example and is not presented as a general accuracy rate.")
    figure(doc, "4.5", SCREEN / "05_compass_verified_response.png", "Compass response grounded in the student's stored record", 6.2)

    heading(doc, "4.3 Adviser implementation", 2)
    body(doc, "The adviser dashboard summarises the assigned level. It presents total students, average CGPA, carryover counts, pass rate, degree class distribution, top performers, students below a risk threshold and recent uploads. Figure 4.6 shows the stable loading state captured while the slower class summary request was still pending. The summary completed later during browser inspection.")
    figure(doc, "4.6", SCREEN / "06_adviser_dashboard.png", "Adviser dashboard while the class summary was loading", 6.2)
    body(doc, "The upload page requires a semester, session and Excel file before preview is enabled. This prevents a file from being written immediately after selection. Figure 4.7 shows the empty form. No new file was confirmed during documentation testing, which avoided changing the live result dataset.")
    figure(doc, "4.7", SCREEN / "07_adviser_upload.png", "Adviser broadsheet upload form", 6.2)
    body(doc, "The history page lists uploads made by the adviser and provides view and delete actions. Figure 4.8 shows the single upload present during the review. The delete action was not tested against the live database because it would remove material project data.")
    figure(doc, "4.8", SCREEN / "08_upload_history.png", "Adviser upload history", 6.2)

    heading(doc, "4.4 Backend implementation", 2)
    body(doc, "FastAPI routers separate the system into student, authentication, upload, result, administrator, analytics, notification and agent endpoints. The health endpoint provides a small availability check. Student endpoints return semester GPA, cumulative GPA and course data. Upload endpoints preview and confirm a workbook. Analytics endpoints compute summaries for the adviser's assigned level. Agent endpoints manage sessions, messages and streamed responses.")
    body(doc, "The result service treats score and unit fields as data that must be validated before calculation. Course codes are normalised so they can match the catalogue. Failed compulsory courses are placed in the outstanding list. A temporal baseline service keeps the cumulative units and points associated with each student and session.")
    body(doc, "The router and service separation reduces repetition. A router reads path, query or form values and returns an HTTP response. A service performs the academic or database work. Shared models describe the expected input. This structure allows the same calculation to be called by a page endpoint or an agent tool without placing the formula in two locations. It also makes an error easier to trace because request handling, calculation and persistence have separate boundaries.")
    body(doc, "The agent route streams response events to the client. Streaming gives the interface early feedback while the model is producing text. Before a record based response is prepared, the agent selects a defined tool and the backend returns a structured observation. The tool result may contain GPA, CGPA, course rows or outstanding courses. Simulation tools accept proposed grades and return calculated values without writing them to the official result tables.")

    heading(doc, "4.5 Database implementation", 2)
    body(doc, "The connected Supabase project was active and hosted in the eu-west-1 region on Postgres 17. Four migration records were present for authentication columns, the administrator table, chat history tables and student session baselines. Realtime publication included students, results and notifications. No storage bucket or Edge Function was present because files are processed by FastAPI and the project uses the database and Auth services directly.")
    body(doc, "The operational schema uses UUID identifiers and foreign keys to connect profiles, courses, uploads and results. Student and adviser profiles store the related Auth user identifier. The application can therefore keep account credentials inside Supabase Auth while storing academic fields in public schema tables. Chat sessions use a matriculation number to group messages. Session baselines store cumulative units and points for a stated academic period.")
    body(doc, "Notifications are created after result processing and published through Supabase Realtime. This supports a responsive student interface without repeated full page refreshes. Realtime delivery does not replace the stored notification row. The row remains the record that determines whether a message exists and whether it has been read. This distinction allows the page to recover its state after a network interruption.")

    heading(doc, "4.6 Test method", 2)
    body(doc, "Testing was carried out on 15 September 2026. The frontend was built with the production build command and checked with Oxlint. The FastAPI and Vite development servers were started locally. Authenticated browser flows were inspected with the supplied student and adviser accounts. Read only SQL and Supabase advisory tools were used to examine table counts, policies and database warnings. API timing used three sequential GET requests for each selected endpoint with a 15 second timeout.")
    body(doc, "The test has limits. It represents one computer, one network path and two accounts. Three timing runs are sufficient to identify a large delay, but they do not replace load testing. The adviser summary was also observed in the browser, where it eventually loaded successfully after multiple parallel requests.")
    body(doc, "Functional inspection followed the main user journeys. The student journey covered sign in, dashboard access, academic record retrieval, Compass history and one stored record question. The adviser journey covered sign in, dashboard loading, the upload form and upload history. The upload confirm and delete actions were not executed because they would change the live dataset. Their presence was inspected through the interface and code, while the report avoids claiming a destructive test that did not occur.")
    body(doc, "The browser viewport was 1,280 by 720 pixels for the captured evidence. Screenshots were taken only after each authenticated page reached a stable state. Detailed class result pages were excluded because they displayed the names and scores of other students. The selected adviser dashboard image records the stable loading state and the adviser account name while the slow summary request was pending. The student images show the project author's test account.")
    body(doc, "Database review used direct counts and consistency queries. It checked the number of records in the main tables, the grade distribution, the active dataset period, duplicate result combinations, orphan references and invalid score or grade values. Supabase security and performance advisers were also read. These checks do not prove that every future workbook is valid, but they describe the state of the connected project at the time of inspection.")

    heading(doc, "4.7 Build and static analysis results", 2)
    table(doc, "4.2", "Build and lint results", ["Test", "Result", "Evidence", "Interpretation"], [
        ["Vite production build", "Pass", "2,304 modules transformed; build completed in 4.64 s", "The frontend produced deployable static assets"],
        ["Bundle size check", "Warning", "Main JavaScript bundle 806.67 kB; gzip 225.58 kB", "Code splitting should be considered"],
        ["Oxlint", "Pass with warnings", "No error; warnings for unused imports, variables and Fast Refresh context", "Warnings do not block the build but should be cleaned"],
    ], [1.3, 1.1, 2.25, 1.65])

    heading(doc, "4.8 API performance results", 2)
    table(doc, "4.3", "API timing results from three sequential runs", ["Endpoint", "HTTP result", "Mean", "Median", "Range", "Finding"], [
        ["GET /health", "200 in 3 of 3", "35.4 ms", "36.7 ms", "17.4 to 52.0 ms", "Fast local availability check"],
        ["GET cumulative GPA", "200 in 3 of 3", "1,046.3 ms", "927.0 ms", "862.7 to 1,349.2 ms", "Acceptable for an interactive record page"],
        ["GET student courses", "200 in 3 of 3", "1,425.9 ms", "1,433.3 ms", "1,388.5 to 1,456.0 ms", "Visible delay but request completed"],
        ["GET adviser summary", "Timeout in 3 of 3", "15,028.5 ms", "15,032.4 ms", "15,009.6 to 15,043.4 ms", "Query path requires optimisation"],
    ], [1.35, 1.0, 0.8, 0.8, 1.2, 1.5])
    body(doc, "The results answer the performance part of the fourth objective with measured evidence. Student retrieval completed in less than 1.5 seconds on average. The adviser summary failed the selected 15 second limit in every isolated run. The dashboard still loaded during browser testing, so the problem is slow completion rather than a permanent functional failure. The route performs several aggregate operations and should be profiled before wider use.")

    heading(doc, "4.9 Functional test results", 2)
    table(doc, "4.4", "Functional test results", ["Test case", "Expected result", "Observed result", "Status"], [
        ["Student sign in", "Valid student reaches student portal", "Success confirmation followed by student dashboard", "Pass"],
        ["Student result view", "Stored result loads for linked profile", "CGPA 4.03, eight courses and outstanding courses displayed", "Pass"],
        ["Compass history", "Saved session and messages load", "Existing session displayed question and grounded answer", "Pass"],
        ["Compass record check", "Answer agrees with result page", "CGPA and two carryovers matched", "Pass for inspected case"],
        ["Adviser sign in", "Active adviser reaches adviser portal", "James Gordon profile opened the dashboard", "Pass"],
        ["Adviser analytics", "Class summary loads", "169 students, average CGPA 3.29 and charts displayed", "Pass with delay"],
        ["Upload form", "Preview disabled until inputs exist", "Semester, session and file were required", "Pass"],
        ["Upload history", "Existing uploads are listed", "One upload with session, row count and date displayed", "Pass"],
        ["Frontend build", "Production assets are generated", "Build completed without error", "Pass"],
    ], [1.45, 2.0, 2.3, 0.85])

    heading(doc, "4.10 Live database results", 2)
    table(doc, "4.5", "Live database evidence at time of inspection", ["Item", "Count", "Meaning"], [
        ["Students", "169", "Student profiles in the application table"],
        ["Advisers", "9", "Adviser profiles across active, pending and revoked states"],
        ["Courses", "37", "Course catalogue records"],
        ["Results", "1,293", "Stored course result rows"],
        ["Uploads", "1", "Confirmed broadsheet upload"],
        ["Notifications", "2,560", "Student notification records"],
        ["Chat sessions", "26", "Saved and temporary conversation headers"],
        ["Chat messages", "84", "Stored user and assistant messages"],
        ["Session baselines", "169", "One current baseline row for each student"],
    ], [1.75, 0.85, 3.8])
    body(doc, "The result table contained 470 A grades, 373 B grades, 346 C grades, 66 D grades and 38 F grades. The distribution sums to 1,293. All current result rows belonged to the 2025/2026 First Semester dataset. Read only consistency checks found no duplicate result combinations, orphan references or invalid score and grade values in the inspected tables.")

    heading(doc, "4.11 Security and database performance review", 2)
    table(doc, "4.6", "Security and performance findings", ["Finding", "Evidence", "Risk", "Required action"], [
        ["Broad RLS policies", "Several main tables permit operations with true conditions", "A client with a public key may access records outside its role", "Replace with ownership and role predicates; revoke excess grants"],
        ["Privileged functions exposed", "Two SECURITY DEFINER functions executable by anon and authenticated roles", "A public RPC path may run elevated code", "Revoke PUBLIC execute or move functions to a private schema"],
        ["Mutable search path", "Two functions lack a fixed search_path", "Object resolution can be unsafe", "Set a safe explicit search_path"],
        ["Leaked password protection off", "Supabase Auth advisor warning", "Known compromised passwords may be accepted", "Enable leaked password protection"],
        ["Unindexed foreign keys", "Five foreign keys lack covering indexes", "Joins and deletes may slow as data grows", "Add indexes after workload review"],
        ["RLS init plan warnings", "Three policies re-evaluate auth functions per row", "Unnecessary policy cost", "Wrap auth calls in a select expression"],
        ["Multiple permissive policies", "Baseline table has overlapping select policies", "Extra evaluation and unclear intent", "Consolidate policies"],
    ], [1.3, 2.0, 1.55, 1.65])
    body(doc, "These findings mean the prototype cannot yet be described as secure for public institutional deployment. The interface separates roles, and the chat tables use user scoped policies, but a determined client can address the Supabase Data API directly. The application should move to least privilege policies before more student records are added. The Supabase security advisor links used in the review point to the database linter, password security and RLS guidance.")

    heading(doc, "4.12 Discussion of results", 2)
    body(doc, "The functional evidence shows that the four core user activities work: authentication, personal result access, adviser analysis and record based conversation. The Compass example also shows why tool grounded advice is useful. The answer repeated values that were visible on the separate result page. This supports the design described by ReAct and tool learning research, where an external observation supplies the factual basis for a response (Yao et al., 2023; Qin et al., 2023).")
    body(doc, "The evaluation also shows the difference between a working demonstration and a deployment ready service. The student endpoints completed within an interactive time range in the small test. The adviser summary did not. The database contained consistent academic rows, yet the access policies were broader than the role model described in Chapter Three. Reporting both outcomes gives a more accurate assessment than relying on screenshots alone.")
    body(doc, "The project objective concerning usability was evaluated through direct task completion and interface inspection. A formal user study was not conducted, so the report does not attach a usability percentage. Future evaluation should recruit students and advisers, record task completion, use a recognised questionnaire such as the System Usability Scale and obtain ethical approval where required.")

    heading(doc, "4.13 Chapter summary", 2)
    body(doc, "The implementation produced a working React and FastAPI application connected to Supabase and Groq. Authenticated screenshots confirm the main student and adviser interfaces. Build, functional, API and database tests provide evidence for the result. The prototype meets its central functions, while the adviser query path and database security controls require further work.")


def chapter_five(doc):
    chapter(doc, "FIVE", "SUMMARY CONCLUSION AND RECOMMENDATIONS")
    heading(doc, "5.1 Summary", 2)
    body(doc, "This project developed Compass to improve the way Computer Science students at Delta State University receive and understand academic results. The existing fallback relies on a shared broadsheet or PDF. Compass keeps the adviser's Excel workflow but moves the processed record into a web application. Each linked student account can view GPA, CGPA, courses and carryovers. Course advisers can upload a broadsheet, inspect history and view class analytics. The advisory agent retrieves stored records through backend tools and explains them in a conversation.")
    body(doc, "The report reviewed learning analytics, large language models, tool using agents and access control. It used ReAct and RBAC as the main frameworks. The design separated the React client, FastAPI services, Supabase data layer and Groq model. Testing covered the production build, linting, authenticated tasks, selected API response times and the live database. The evidence confirmed working student and adviser flows. It also identified a slow adviser summary and important RLS and function permission weaknesses.")

    heading(doc, "5.2 Achievement of objectives", 2)
    table(doc, "5.1", "Achievement of project objectives", ["Objective", "Evidence", "Assessment"], [
        ["Design a role based web platform", "Separate student, adviser and administrator routes; authenticated student and adviser screenshots", "Functionally achieved; database policies need hardening"],
        ["Develop broadsheet ingestion and anomaly checking", "Preview and confirmation routes, Excel parsing services, grade and baseline checks", "Achieved for supported workbook layouts"],
        ["Implement a tool using academic adviser", "Compass sessions, academic tools and a response matching the stored result", "Achieved for the inspected record based case"],
        ["Evaluate the system", "Build, lint, functional, timing, consistency and Supabase advisor results", "Achieved within the stated small test scope"],
    ], [2.05, 3.0, 1.4])

    heading(doc, "5.3 Contributions of the project", 2)
    body(doc, "The first contribution is an integrated departmental workflow that converts an existing Excel broadsheet into individual web records. The second is a transparent advisory design in which calculations and retrieval remain in backend tools. The third is a temporal baseline mechanism that preserves cumulative totals when a new semester is imported. The fourth is an evidence based evaluation that records successful functions and remaining risks.")

    heading(doc, "5.4 Conclusion", 2)
    body(doc, "Compass shows that a departmental result system can provide more useful access than a shared static file. Students can find their own record quickly, track outstanding courses and ask a question in familiar language. Advisers can work from their existing Excel input and receive a class summary. The advisory agent is useful because it reads database values through defined tools before composing a response.")
    body(doc, "The implementation is a working prototype. It should not hold a wider institutional dataset until access policies are corrected and the adviser analytics route is optimised. With those controls in place, the architecture can support a larger pilot and a formal usability study.")

    heading(doc, "5.5 Recommendations", 2)
    numbered_list(doc, [
        "Replace permissive public table policies with explicit ownership and role checks based on auth.uid(), and revoke table privileges that are not required.",
        "Revoke public execution of privileged functions, move internal functions to a private schema where possible and set an explicit safe search path.",
        "Add covering indexes for the foreign keys identified by the Supabase performance advisor after checking expected query patterns.",
        "Profile the adviser dashboard summary, reduce repeated queries and consider one server side aggregate query or cached summary for the selected session.",
        "Split the main frontend bundle with route based lazy loading so the first download is smaller.",
        "Add automated unit and integration tests for grade boundaries, GPA calculations, workbook layouts, role checks and agent tools.",
        "Conduct a supervised pilot with students and course advisers, using task completion measures and a standard usability questionnaire.",
        "Prepare a written data retention policy for result files, notifications and chat history, with a clear process for correcting inaccurate data.",
    ])

    heading(doc, "5.6 Limitations", 2)
    body(doc, "The evaluation used one student account and one adviser account. It did not include an administrator screenshot or a multiuser load test. The live result dataset represented one semester, so long term behaviour across several uploads was not observed. External network services affected response time. The study did not measure learning outcomes, adviser workload reduction or student satisfaction through a controlled experiment.")

    heading(doc, "5.7 Suggestions for future work", 2)
    body(doc, "Future work should add a policy hardened staging environment and automated security tests before another data import. A broader pilot can then study whether students understand their academic standing more accurately after using Compass. The system can also add degree requirement rules approved by the department, accessible audit logs, adviser comments and export of a personal transcript summary. Any predictive model should be introduced only with a documented dataset, fairness review and separate evaluation from the official result calculations.")


def add_references(doc):
    heading(doc, "REFERENCES", 1)
    refs = [
        "Albreiki, B., Zaki, N., & Alashwal, H. (2021). A systematic literature review of student performance prediction using machine learning techniques. Education Sciences, 11(9), 552. https://doi.org/10.3390/educsci11090552",
        "Beck, K., Beedle, M., van Bennekum, A., Cockburn, A., Cunningham, W., Fowler, M., Grenning, J., Highsmith, J., Hunt, A., Jeffries, R., Kern, J., Marick, B., Martin, R. C., Mellor, S., Schwaber, K., Sutherland, J., & Thomas, D. (2001). Manifesto for Agile software development. https://agilemanifesto.org/",
        "Bender, E. M., Gebru, T., McMillan-Major, A., & Shmitchell, S. (2021). On the dangers of stochastic parrots: Can language models be too big? Proceedings of the 2021 ACM Conference on Fairness, Accountability, and Transparency, 610-623. https://doi.org/10.1145/3442188.3445922",
        "Bommasani, R., Hudson, D. A., Adeli, E., Altman, R., Arora, S., von Arx, S., Bernstein, M. S., Bohg, J., Bosselut, A., Brunskill, E., Brynjolfsson, E., Buch, S., Card, D., Castellon, R., Chatterji, N., Chen, A., Creel, K., Davis, J. Q., Demszky, D., ... Liang, P. (2021). On the opportunities and risks of foundation models. arXiv. https://doi.org/10.48550/arXiv.2108.07258",
        "Brown, T. B., Mann, B., Ryder, N., Subbiah, M., Kaplan, J., Dhariwal, P., Neelakantan, A., Shyam, P., Sastry, G., Askell, A., Agarwal, S., Herbert-Voss, A., Krueger, G., Henighan, T., Child, R., Ramesh, A., Ziegler, D. M., Wu, J., Winter, C., ... Amodei, D. (2020). Language models are few-shot learners. Advances in Neural Information Processing Systems, 33, 1877-1901.",
        "Devlin, J., Chang, M. W., Lee, K., & Toutanova, K. (2019). BERT: Pre-training of deep bidirectional transformers for language understanding. Proceedings of NAACL-HLT 2019, 4171-4186. https://doi.org/10.18653/v1/N19-1423",
        "Ferraiolo, D. F., Sandhu, R., Gavrila, S., Kuhn, D. R., & Chandramouli, R. (2001). Proposed NIST standard for role-based access control. ACM Transactions on Information and System Security, 4(3), 224-274. https://doi.org/10.1145/501978.501980",
        "Fielding, R. T. (2000). Architectural styles and the design of network-based software architectures [Doctoral dissertation, University of California, Irvine].",
        "International Organization for Standardization. (2023). ISO/IEC 25010:2023 systems and software engineering: Systems and software quality models.",
        "Ji, Z., Lee, N., Frieske, R., Yu, T., Su, D., Xu, Y., Ishii, E., Bang, Y. J., Madotto, A., & Fung, P. (2023). Survey of hallucination in natural language generation. ACM Computing Surveys, 55(12), Article 248. https://doi.org/10.1145/3571730",
        "Kabathova, J., & Drlik, M. (2021). Towards predicting student's dropout in university courses using different machine learning techniques. Applied Sciences, 11(7), 3130. https://doi.org/10.3390/app11073130",
        "Kasneci, E., Sessler, K., Küchemann, S., Bannert, M., Dementieva, D., Fischer, F., Gasser, U., Groh, G., Günnemann, S., Hüllermeier, E., Krusche, S., Kutyniok, G., Michaeli, T., Nerdel, C., Pfeffer, J., Poquet, O., Sailer, M., Schmidt, A., Seidel, T., ... Kasneci, G. (2023). ChatGPT for good? On opportunities and challenges of large language models for education. Learning and Individual Differences, 103, 102274. https://doi.org/10.1016/j.lindif.2023.102274",
        "Khosravi, H., Shum, S. B., Chen, G., Conati, C., Tsai, Y. S., Kay, J., Knight, S., Martinez-Maldonado, R., Sadiq, S., & Gašević, D. (2022). Explainable artificial intelligence in education. Computers and Education: Artificial Intelligence, 3, 100074. https://doi.org/10.1016/j.caeai.2022.100074",
        "Liu, X., Yu, H., Zhang, H., Xu, Y., Lei, X., Lai, H., Gu, Y., Ding, H., Men, K., Yang, K., Zhang, S., Deng, X., Zeng, A., Du, Z., Zhang, C., Shen, S., Zhang, T., Su, Y., Sun, H., ... Tang, J. (2024). AgentBench: Evaluating LLMs as agents. International Conference on Learning Representations. https://doi.org/10.48550/arXiv.2308.03688",
        "Lo, C. K. (2023). What is the impact of ChatGPT on education? A rapid review of the literature. Education Sciences, 13(4), 410. https://doi.org/10.3390/educsci13040410",
        "Namoun, A., & Alshanqiti, A. (2021). Predicting student performance using data mining and learning analytics techniques: A systematic literature review. Applied Sciences, 11(1), 237. https://doi.org/10.3390/app11010237",
        "National Institute of Standards and Technology. (2023). Artificial intelligence risk management framework (AI RMF 1.0). https://doi.org/10.6028/NIST.AI.100-1",
        "Nielsen, J. (1994). Usability engineering. Morgan Kaufmann.",
        "Nigeria Data Protection Act, 2023.",
        "Ouyang, F., Zheng, L., & Jiao, P. (2022). Artificial intelligence in online higher education: A systematic review of empirical research from 2011 to 2020. Education and Information Technologies, 27, 7893-7925. https://doi.org/10.1007/s10639-022-10925-9",
        "Park, J. S., O'Brien, J. C., Cai, C. J., Morris, M. R., Liang, P., & Bernstein, M. S. (2023). Generative agents: Interactive simulacra of human behavior. Proceedings of the 36th Annual ACM Symposium on User Interface Software and Technology, Article 2. https://doi.org/10.1145/3586183.3606763",
        "Qin, Y., Liang, S., Ye, Y., Zhu, K., Yan, L., Lu, Y., Lin, Y., Cong, X., Tang, X., Qian, B., Zhao, S., Tian, R., Xie, R., Zhou, J., Gerstein, M., Li, D., Liu, Z., & Sun, M. (2023). Tool learning with foundation models. arXiv. https://doi.org/10.48550/arXiv.2304.08354",
        "Sandhu, R. S., Coyne, E. J., Feinstein, H. L., & Youman, C. E. (1996). Role-based access control models. Computer, 29(2), 38-47. https://doi.org/10.1109/2.485845",
        "Schick, T., Dwivedi-Yu, J., Dessì, R., Raileanu, R., Lomeli, M., Zettlemoyer, L., Cancedda, N., & Scialom, T. (2023). Toolformer: Language models can teach themselves to use tools. Advances in Neural Information Processing Systems, 36, 68539-68551.",
        "Supabase. (2026). Row level security. https://supabase.com/docs/guides/database/postgres/row-level-security",
        "Tlili, A., Shehata, B., Adarkwah, M. A., Bozkurt, A., Hickey, D. T., Huang, R., & Agyemang, B. (2023). What if the devil is my guardian angel? ChatGPT as a case study of using chatbots in education. Smart Learning Environments, 10, 15. https://doi.org/10.1186/s40561-023-00237-x",
        "UNESCO. (2023). Guidance for generative AI in education and research. UNESCO.",
        "Vaswani, A., Shazeer, N., Parmar, N., Uszkoreit, J., Jones, L., Gomez, A. N., Kaiser, Ł., & Polosukhin, I. (2017). Attention is all you need. Advances in Neural Information Processing Systems, 30, 5998-6008.",
        "Wei, J., Wang, X., Schuurmans, D., Bosma, M., Xia, F., Chi, E., Le, Q. V., & Zhou, D. (2022). Chain-of-thought prompting elicits reasoning in large language models. Advances in Neural Information Processing Systems, 35, 24824-24837.",
        "Wolf, T., Debut, L., Sanh, V., Chaumond, J., Delangue, C., Moi, A., Cistac, P., Rault, T., Louf, R., Funtowicz, M., Davison, J., Shleifer, S., von Platen, P., Ma, C., Jernite, Y., Plu, J., Xu, C., Le Scao, T., Gugger, S., ... Rush, A. M. (2020). Transformers: State-of-the-art natural language processing. Proceedings of EMNLP 2020: System Demonstrations, 38-45. https://doi.org/10.18653/v1/2020.emnlp-demos.6",
        "Yao, S., Zhao, J., Yu, D., Du, N., Shafran, I., Narasimhan, K., & Cao, Y. (2023). ReAct: Synergizing reasoning and acting in language models. International Conference on Learning Representations. https://doi.org/10.48550/arXiv.2210.03629",
        "Zawacki-Richter, O., Marín, V. I., Bond, M., & Gouverneur, F. (2019). Systematic review of research on artificial intelligence applications in higher education. International Journal of Educational Technology in Higher Education, 16, 39. https://doi.org/10.1186/s41239-019-0171-0",
        "Zhao, W. X., Zhou, K., Li, J., Tang, T., Wang, X., Hou, Y., Min, Y., Zhang, B., Zhang, J., Dong, Z., Du, Y., Yang, C., Chen, Y., Chen, Z., Jiang, J., Ren, R., Li, Y., Tang, X., Liu, Z., ... Wen, J. R. (2023). A survey of large language models. arXiv. https://doi.org/10.48550/arXiv.2303.18223",
    ]
    for ref in refs:
        p = body(doc, ref, indent=False)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.hanging_indent = Inches(0.5)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(8)


def add_appendices(doc):
    heading(doc, "APPENDICES", 1)
    heading(doc, "Appendix A Selected API routes", 2)
    table(doc, "A.1", "Selected FastAPI routes", ["Method", "Route", "Purpose"], [
        ["GET", "/health", "Check server availability"],
        ["GET", "/students/{matric}/gpa/cumulative", "Return cumulative academic totals"],
        ["GET", "/students/{matric}/courses", "Return a student's course results"],
        ["POST", "/upload/preview", "Parse and validate a broadsheet without confirmation"],
        ["POST", "/upload/confirm", "Persist a reviewed upload"],
        ["GET", "/analytics/dashboard-summary", "Return adviser class summary"],
        ["POST", "/agent/chat/stream", "Stream a Compass response"],
        ["GET", "/agent/sessions/{matric}", "List a student's chat sessions"],
    ], [0.75, 3.0, 2.7])

    heading(doc, "Appendix B Grade point rules", 2)
    table(doc, "B.1", "Five point grade mapping used by the prototype", ["Score range", "Grade", "Grade point"], [
        ["70 to 100", "A", "5"], ["60 to 69", "B", "4"], ["50 to 59", "C", "3"],
        ["45 to 49", "D", "2"], ["0 to 44", "F", "0"],
    ], [2.2, 1.5, 2.1])

    heading(doc, "Appendix C Test evidence summary", 2)
    table(doc, "C.1", "Test evidence summary", ["Item", "Recorded evidence"], [
        ["Test date", "15 September 2026"],
        ["Frontend build", "Passed"],
        ["Frontend lint", "Completed with warnings and no errors"],
        ["Authenticated roles", "Student and active adviser"],
        ["Student API timing", "Three runs per selected endpoint"],
        ["Adviser summary timing", "Three runs, each reached the 15 second timeout"],
        ["Database access", "Read only during documentation"],
        ["Destructive actions", "None"],
    ], [2.2, 4.1])

    heading(doc, "Appendix D Ethical and AI use statement", 2)
    body(doc, "The project processes academic records and therefore requires limited disclosure, controlled access and accurate correction procedures. Screenshots selected for this report avoid detailed class result pages that display the names and scores of other students. The documentation process used an AI writing assistant for source organisation, language editing, diagram generation and document formatting. The author remains responsible for verifying the implementation, references, test evidence and final submission. AI generated statements were not treated as test results. Database counts, interface behaviour and response timings were obtained from the working project and connected Supabase instance.")


def build():
    doc = Document()
    setup_styles(doc)
    add_preliminaries(doc)

    main_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main_section)
    main_section.header.is_linked_to_previous = False
    main_section.footer.is_linked_to_previous = False
    set_page_number_start(main_section, 1, "decimal")
    header_p = main_section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(header_p, roman=False)
    main_section.footer.paragraphs[0].text = ""

    chapter_one(doc)
    chapter_two(doc)
    chapter_three(doc)
    chapter_four(doc)
    chapter_five(doc)
    add_references(doc)
    add_appendices(doc)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    props = doc.core_properties
    props.title = TITLE.title()
    props.subject = "Bachelor of Science final year project report"
    props.author = AUTHOR.title()
    props.keywords = "academic results, learning analytics, agentic AI, Supabase, FastAPI, React"

    DOC_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
